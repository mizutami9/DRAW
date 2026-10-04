import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\mizut\Downloads\NICO_DRAW_creator_outreach_group_leaders_v200.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

ws_main = wb["独立営業先_メイン"]
print(f"独立営業先_メイン rows: {ws_main.max_row} (Data rows: {ws_main.max_row - 1})")

ws_group = wb["グループメンバー一覧"]
print(f"グループメンバー一覧 rows: {ws_group.max_row} (Data rows: {ws_group.max_row - 1})")

ws_ex = wb["統合・除外候補一覧"]
print(f"統合・除外候補一覧 rows: {ws_ex.max_row} (Data rows: {ws_ex.max_row - 1})")

ws_dash = wb["集計ダッシュボード"]
print(f"集計ダッシュボード rows: {ws_dash.max_row}")

# Check sample data in main
print("\nSample 3 entries from main:")
for r in range(2, 5):
    vals = [ws_main.cell(r, c).value for c in [1, 2, 3, 4, 5, 6]]
    print(" ", vals)

print("\nSample 3 new entries from main:")
for r in range(ws_main.max_row - 2, ws_main.max_row + 1):
    vals = [ws_main.cell(r, c).value for c in [1, 2, 3, 4, 5, 6]]
    print(" ", vals)
