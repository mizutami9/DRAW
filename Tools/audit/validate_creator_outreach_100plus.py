from __future__ import annotations

from collections import Counter
from pathlib import Path
import json
import re
import sys
import zipfile

import openpyxl


sys.stdout.reconfigure(encoding="utf-8")
FILE = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_FINAL_100PLUS.xlsx")
EXPECTED_SHEETS = ["最終営業リスト", "A_最優先", "B_営業対象", "保留", "除外", "グループ統合監査", "地域別進捗", "DMテンプレ", "調査ログ"]
EXPECTED_HEADERS = [
    "No.", "優先度", "営業先名", "営業単位", "営業言語", "地域", "X URL", "X本人確認",
    "Xフォロワー", "YouTube URL", "YouTube登録者", "Twitch URL", "Twitchフォロワー",
    "主な連絡方法", "DM可否", "代替連絡方法", "複数人プレイ確認", "共演者",
    "固定/準固定/単発", "確認ゲーム", "証拠URL1", "証拠URL2", "証拠URL3", "最終活動日",
    "確認日", "NICO DRAW適合理由", "送信済", "送信日", "返信", "デモプレイ",
    "配信/動画化", "製品版キー予定人数", "備考",
]


def clean_url(value):
    value = str(value or "").strip()
    if not value.startswith("http"):
        return ""
    return value.split("?", 1)[0].rstrip("/").replace("twitter.com/", "x.com/").replace("www.x.com/", "x.com/").lower()


def name_key(value):
    return re.sub(r"[^a-z0-9\u3040-\u30ff\u3400-\u9fff\uac00-\ud7af]+", "", str(value or "").lower())


def rows(sheet):
    headers = [c.value for c in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in row)]


def main():
    problems = []
    with zipfile.ZipFile(FILE) as archive:
        bad_member = archive.testzip()
        if bad_member:
            problems.append(f"ZIP破損: {bad_member}")
    book = openpyxl.load_workbook(FILE, data_only=False)
    if book.sheetnames != EXPECTED_SHEETS:
        problems.append(f"シート順不一致: {book.sheetnames}")
    final = book["最終営業リスト"]
    headers = [c.value for c in final[1]]
    if headers != EXPECTED_HEADERS:
        problems.append("最終営業リストの列構成不一致")
    data = rows(final)
    if not 120 <= len(data) <= 150:
        problems.append(f"最終件数が120〜150外: {len(data)}")
    if [r["No."] for r in data] != list(range(1, len(data) + 1)):
        problems.append("No.が連番でない")
    if any(r["優先度"] not in {"A", "B"} for r in data):
        problems.append("A/B以外が最終リストに混入")

    required = ["営業先名", "営業単位", "営業言語", "地域", "主な連絡方法", "DM可否", "複数人プレイ確認", "確認ゲーム", "証拠URL1", "最終活動日", "確認日", "NICO DRAW適合理由"]
    for idx, row in enumerate(data, 2):
        for field in required:
            if row.get(field) in (None, ""):
                problems.append(f"行{idx}: {field}空欄")
        if not str(row["証拠URL1"]).startswith("http"):
            problems.append(f"行{idx}: 直接証拠URLなし")
        if not re.fullmatch(r"20\d{2}-\d{2}-\d{2}", str(row["最終活動日"])):
            problems.append(f"行{idx}: 最終活動日形式不正")
        if not str(row["主な連絡方法"]).find("http") >= 0 and not str(row["X URL"]).startswith("http"):
            problems.append(f"行{idx}: 実URLを伴う主連絡経路なし")

    names = Counter(name_key(r["営業先名"]) for r in data)
    dup_names = [name for name, count in names.items() if count > 1]
    if dup_names:
        problems.append(f"営業先名重複: {dup_names}")
    for field in ("X URL", "YouTube URL", "Twitch URL"):
        urls = Counter(clean_url(r[field]) for r in data if clean_url(r[field]))
        dup_urls = [url for url, count in urls.items() if count > 1]
        if dup_urls:
            problems.append(f"{field}重複: {dup_urls}")

    a_rows = rows(book["A_最優先"])
    b_rows = rows(book["B_営業対象"])
    if len(a_rows) != sum(r["優先度"] == "A" for r in data):
        problems.append("Aシート件数不一致")
    if len(b_rows) != sum(r["優先度"] == "B" for r in data):
        problems.append("Bシート件数不一致")
    if {name_key(r["営業先名"]) for r in a_rows + b_rows} != {name_key(r["営業先名"]) for r in data}:
        problems.append("A/Bシートと最終リストの集合不一致")
    if any(r["優先度"] != "A" for r in a_rows) or any(r["優先度"] != "B" for r in b_rows):
        problems.append("A/Bシートに誤優先度が混入")

    hold_names = {name_key(r.get("候補名")) for r in rows(book["保留"])}
    excluded_names = {name_key(r.get("候補名")) for r in rows(book["除外"])}
    final_names = set(names)
    overlaps = {
        "最終×保留": sorted(final_names & hold_names),
        "最終×除外": sorted(final_names & excluded_names),
    }
    for label, items in overlaps.items():
        if items:
            problems.append(f"{label}重複: {items[:10]}")

    dm_languages = {str(r.get("言語")) for r in rows(book["DMテンプレ"])}
    final_languages = {str(r["営業言語"]) for r in data}
    missing_templates = sorted(final_languages - dm_languages)
    if missing_templates:
        problems.append(f"DMテンプレ不足: {missing_templates}")

    # Verify real Excel hyperlinks in all primary URL cells.
    for sheet_name in ("最終営業リスト", "A_最優先", "B_営業対象"):
        sheet = book[sheet_name]
        hs = [c.value for c in sheet[1]]
        for field in ("X URL", "YouTube URL", "Twitch URL", "証拠URL1", "証拠URL2", "証拠URL3"):
            col = hs.index(field) + 1
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, col)
                if str(cell.value or "").startswith("http") and not cell.hyperlink:
                    problems.append(f"{sheet_name}!{cell.coordinate}: ハイパーリンク未設定")

    if len(final.data_validations.dataValidation) < 5:
        problems.append("進捗入力規則が不足")
    formulas = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            formulas.extend(str(c.value) for c in row if isinstance(c.value, str) and c.value.startswith("="))

    result = {
        "file": str(FILE), "bytes": FILE.stat().st_size, "sheets": book.sheetnames,
        "final": len(data), "A": len(a_rows), "B": len(b_rows), "hold": len(rows(book["保留"])),
        "excluded": len(rows(book["除外"])), "audit_log": len(rows(book["調査ログ"])),
        "languages": Counter(r["営業言語"] for r in data),
        "x_verified": sum(str(r["X URL"]).startswith("http") for r in data),
        "non_x_contact": sum(not str(r["X URL"]).startswith("http") for r in data),
        "formulas": len(formulas), "validations": len(final.data_validations.dataValidation),
        "problems": problems,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if problems:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
