from collections import Counter, defaultdict
from datetime import date
from difflib import SequenceMatcher
import json
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
TODAY = date(2026, 10, 3)
ACTIVE_DAYS = 31


def canonical_x(value):
    match = re.search(r"(?:x|twitter)\.com/([A-Za-z0-9_]+)", str(value or ""), re.I)
    return f"https://x.com/{match.group(1)}" if match else None


def canonical_url(value):
    return str(value or "").strip().split("?")[0].rstrip("/").lower()


wb = openpyxl.load_workbook("NICO_DRAW_creator_outreach_100PLUS_WORKING.xlsx", data_only=True)
existing_names, existing_x, existing_youtube = set(), set(), set()
for sheet_name in ("最終営業リスト", "保留", "除外"):
    ws = wb[sheet_name]
    headers = {cell.value: cell.column for cell in ws[1]}
    name_col = headers.get("営業先名") or headers.get("候補名")
    x_col = headers.get("X URL(本人確認済のみ)") or headers.get("X URL(確認済のみ)")
    yt_col = headers.get("YouTube URL") or headers.get("YouTube")
    for row in range(2, ws.max_row + 1):
        if name_col:
            existing_names.add(re.sub(r"\W", "", str(ws.cell(row, name_col).value or "").lower()))
        if x_col:
            x = canonical_x(ws.cell(row, x_col).value)
            if x:
                existing_x.add(x.lower())
        if yt_col:
            value = canonical_url(ws.cell(row, yt_col).value)
            if value.startswith("http"):
                existing_youtube.add(value)

data = json.load(open("Tools/audit/youtube_discovery.json", encoding="utf-8"))
eligible, rejected = [], []
for candidate in data["candidates"]:
    reasons = []
    last = date.fromisoformat(candidate["last_active"])
    if (TODAY - last).days > ACTIVE_DAYS:
        reasons.append("31日超")
    x = canonical_x(candidate.get("x_url"))
    youtube = canonical_url(candidate["youtube_url"])
    name = re.sub(r"\W", "", candidate["name"].lower())
    if x and x.lower() in existing_x:
        reasons.append("既存X")
    if youtube in existing_youtube:
        reasons.append("既存YouTube")
    if name in existing_names:
        reasons.append("既存名")
    if not candidate.get("proofs"):
        reasons.append("証拠なし")
    if not candidate.get("external_links"):
        reasons.append("連絡先なし")
    candidate = {**candidate, "x_url": x}
    if reasons:
        rejected.append({"name": candidate["name"], "reasons": reasons})
    else:
        score = (
            min(len(candidate["proofs"]), 3) * 3
            + min(len(candidate["games"]), 3) * 2
            + (2 if x else 0)
            + (2 if (TODAY - last).days <= 14 else 0)
            + (1 if candidate.get("subscribers") is not None and candidate["subscribers"] <= 15_000 else 0)
        )
        candidate["score"] = score
        eligible.append(candidate)

# Xなしでも公式AboutのInstagram/Bluesky/リンク集等があるため営業可能だが、Xありを優先。
eligible.sort(key=lambda c: (-c["score"], not bool(c["x_url"]), c.get("subscribers") is None, c.get("subscribers") or 999999, c["name"].lower()))

print(json.dumps({
    "eligible": len(eligible),
    "eligible_language": Counter(c["language"] for c in eligible),
    "with_x": sum(bool(c["x_url"]) for c in eligible),
    "proof_counts": Counter(len(c["proofs"]) for c in eligible),
    "games_counts": Counter(len(c["games"]) for c in eligible),
    "rejected": rejected,
    "candidates": [{
        "name": c["name"], "language": c["language"], "score": c["score"],
        "subs": c["subscribers"], "x": c["x_url"], "last": c["last_active"],
        "games": c["games"], "proofs": len(c["proofs"]), "youtube": c["youtube_url"],
    } for c in eligible],
}, ensure_ascii=False, indent=2, default=dict))
