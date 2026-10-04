from pathlib import Path
import json
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
data = json.loads(Path("Tools/audit/creator_evidence_refresh.json").read_text(encoding="utf-8"))["channels"]
workbook = openpyxl.load_workbook("NICO_DRAW_creator_outreach_WORKING.xlsx")
sheet = workbook["最終営業リスト"]
headers = {cell.value: cell.column for cell in sheet[1]}

mismatches = []
generic = 0
channel_count = 0
video_count = 0
with_socials = 0
for row in range(2, sheet.max_row + 1):
    twitch = str(sheet.cell(row, headers["Twitch URL"]).value or "")
    match = re.search(r"twitch\.tv/([A-Za-z0-9_]+)", twitch, re.I)
    x_url = str(sheet.cell(row, headers["X URL(本人確認済のみ)"]).value or "")
    if match:
        channel_count += 1
        channel = data[match.group(1).lower()]
        video_count += len(channel.get("videos") or [])
        social_urls = " ".join(
            social.get("url", "")
            for social in ((channel.get("channel") or {}).get("socialMedias") or [])
        )
        with_socials += bool(social_urls)
        if x_url.startswith("http"):
            handle = re.search(r"(?:x|twitter)\.com/([A-Za-z0-9_]+)", x_url, re.I)
            linked = {
                m.group(1).lower()
                for m in re.finditer(r"(?:x|twitter)\.com/([A-Za-z0-9_]+)", social_urls, re.I)
            }
            if handle and handle.group(1).lower() not in linked:
                mismatches.append({
                    "row": row,
                    "name": sheet.cell(row, headers["営業先名"]).value,
                    "x": x_url,
                    "twitch_socials": social_urls,
                })
    if "/videos" in str(sheet.cell(row, headers["証拠URL"]).value or ""):
        generic += 1

print(json.dumps({
    "final_twitch_channels": channel_count,
    "retrieved_videos": video_count,
    "channels_with_social_links": with_socials,
    "x_crosslink_mismatches": mismatches,
    "generic_evidence_rows": generic,
}, ensure_ascii=False, indent=2))
