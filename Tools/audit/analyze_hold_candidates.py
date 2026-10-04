from collections import Counter
import json
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
wb = openpyxl.load_workbook("NICO_DRAW_creator_outreach_100PLUS_WORKING.xlsx", data_only=True)
ws = wb["保留"]
headers = {cell.value: cell.column for cell in ws[1]}

reasons = Counter()
near = []
for row in range(2, ws.max_row + 1):
    reason = str(ws.cell(row, headers["理由"]).value or "")
    if "複数人プレイを確認できず" in reason:
        reasons["multiplayer_missing"] += 1
    elif "最終活動日を確認できない" in reason:
        reasons["activity_missing"] += 1
    elif "営業導線" in reason or "連絡" in reason:
        reasons["contact_missing"] += 1
    else:
        reasons[reason[:80] or "other"] += 1

    multi_count = ws.cell(row, headers["複数人プレイ確認数"]).value
    friend_count = ws.cell(row, headers["友人/共演系確認数"]).value
    last = ws.cell(row, headers["最終活動日"]).value
    x_url = ws.cell(row, headers["X URL(確認済のみ)"]).value
    twitch = ws.cell(row, headers["Twitch"]).value
    youtube = ws.cell(row, headers["YouTube"]).value
    if (multi_count or 0) > 0 or (friend_count or 0) > 0:
        near.append({
            "row": row,
            "name": ws.cell(row, headers["候補名"]).value,
            "reason": reason,
            "last": last,
            "x": x_url,
            "twitch": twitch,
            "youtube": youtube,
            "multi": multi_count,
            "friend": friend_count,
        })

print(json.dumps({"reason_counts": reasons, "near_count": len(near), "near": near}, ensure_ascii=False, indent=2))
