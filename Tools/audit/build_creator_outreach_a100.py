from __future__ import annotations

from collections import Counter
from copy import copy
from pathlib import Path
import json
import re
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from a100_config import DM_DONE_NAMES, is_dm_done, normalize


sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"C:\workSpace\PICO")
SOURCE = ROOT / "NICO_DRAW_creator_outreach_FINAL_100PLUS.xlsx"
SCORED = ROOT / "Tools" / "audit" / "a100_scored.json"
OUTPUT = ROOT / "NICO_DRAW_creator_outreach_FINAL_A100.xlsx"
CHECKED = "2026-10-04"

FINAL_HEADERS = [
    "No.", "優先度", "営業先名", "営業単位", "営業言語", "地域", "X URL", "X本人確認",
    "Xフォロワー", "YouTube URL", "YouTube登録者", "Twitch URL", "Twitchフォロワー",
    "主な連絡方法", "DM可否", "代替連絡方法", "複数人プレイ確認", "共演者",
    "固定/準固定/単発", "確認ゲーム", "証拠URL1", "証拠URL2", "証拠URL3", "最終活動日",
    "確認日", "NICO DRAW適合理由", "送信済", "送信日", "返信", "デモプレイ",
    "配信/動画化", "製品版キー予定人数", "備考",
]

# Exactly the 61 already-contacted independent contacts provided by the user.
DM_DONE_RECORDS = [
    ("キュレ", "tmpCure@mail.com", "日本語", "日本"),
    ("にらの民", "https://x.com/NiranoTami / nirany133nira@outlook.jp", "日本語", "日本"),
    ("あんぜん☆だいいち", "https://x.com/anzen_daiichi_4", "日本語", "日本"),
    ("平方ワールド", "https://x.com/HeihouWorld", "日本語", "日本"),
    ("カタリベゲームズ", "https://x.com/nii_san_213", "日本語", "日本"),
    ("ショウセイshousei-agent", "https://x.com/shousei_agent", "日本語", "日本"),
    ("サ柊", "https://x.com/shu_of_service", "日本語", "日本"),
    ("ヨフカシ@ゲーム実況", "https://x.com/yohu_kashi", "日本語", "日本"),
    ("shooon＠ゲーム実況", "https://x.com/k14440", "日本語", "日本"),
    ("べんべん", "https://www.youtube.com/watch?v=U38Z5WSkH_Y", "日本語", "日本"),
    ("かねぶろ", "https://www.youtube.com/@QuarBro/videos", "日本語", "日本"),
    ("毎日がエブリデイライフ", "https://www.youtube.com/@毎日がエブリデイライフ/featured", "日本語", "日本"),
    ("さっぽこぽんちゃい", "https://www.youtube.com/@sappoko_ponchai", "日本語", "日本"),
    ("とまとなすびチャンネル", "https://www.youtube.com/@tomatonasubich/videos", "日本語", "日本"),
    ("TMG", "https://www.youtube.com/@TMG-4manzi/videos", "日本語", "日本"),
    ("KAFE", "https://www.youtube.com/@kafecafe/videos", "日本語", "日本"),
    ("夜な夜なドロン", "https://www.youtube.com/@yonayonadoron/videos", "日本語", "日本"),
    ("ぴきねこ", "https://www.youtube.com/@pikineco/videos", "日本語", "日本"),
    ("ゆるゆるなまな板さん家", "https://www.youtube.com/@manaitasanchi/streams", "日本語", "日本"),
    ("しょぴ", "https://www.youtube.com/@sho-p/streams", "日本語", "日本"),
    ("怀特尼斯whiteness", "https://space.bilibili.com/92682962/upload/video", "中国語", "中国"),
    ("折笔日记", "https://space.bilibili.com/903630/upload/video", "中国語", "中国"),
    ("代糖L", "https://www.bilibili.com/video/BV14Hgy6AECL/", "中国語", "中国"),
    ("扶-城", "https://space.bilibili.com/1371735722", "中国語", "中国"),
    ("紫幻青", "https://space.bilibili.com/170805608", "中国語", "中国"),
    ("無畏寒東", "https://space.bilibili.com/3494371580184942", "中国語", "中国"),
    ("王大大Gamer", "https://space.bilibili.com/85400122/", "中国語", "中国"),
    ("Emberwindy", "https://space.bilibili.com/1455953", "中国語", "中国"),
    ("蓝若羽", "https://space.bilibili.com/228803159", "中国語", "中国"),
    ("水天葵", "https://space.bilibili.com/18306130", "中国語", "中国"),
    ("哇系小黑哥", "https://space.bilibili.com/150799", "中国語", "中国"),
    ("远哥视频", "https://space.bilibili.com/15152799", "中国語", "中国"),
    ("爱打游戏的老给", "https://space.bilibili.com/73652352", "中国語", "中国"),
    ("龙眼皮Lyd", "https://space.bilibili.com/39846961", "中国語", "中国"),
    ("清照儿", "https://space.bilibili.com/737453", "中国語", "中国"),
    ("Zero零八", "https://space.bilibili.com/3707040528599288", "中国語", "中国"),
    ("梧箐箐", "https://space.bilibili.com/1629955", "中国語", "中国"),
    ("老白ya-游戏", "https://space.bilibili.com/32048607", "中国語", "中国"),
    ("玲珑酱酱子", "https://space.bilibili.com/14060509", "中国語", "中国"),
    ("沐锶MuSr", "https://space.bilibili.com/2027700528", "中国語", "中国"),
    ("云端不是云玩家", "https://space.bilibili.com/3546780199291598", "中国語", "中国"),
    ("乱打乱蒙", "https://space.bilibili.com/18982459", "中国語", "中国"),
    ("Lominoth", "https://x.com/Lominoth", "英語", "米国（ユーザー指定）"),
    ("Kino & Lotus | ENVtubers", "https://x.com/KinaurandLobus", "英語", "米国（ユーザー指定）"),
    ("Elie", "https://x.com/itseliep", "英語", "米国（ユーザー指定）"),
    ("Zoey The Demon Fox", "https://x.com/ZoeyTheDemonFox", "英語", "米国（ユーザー指定）"),
    ("OniHeart", "https://x.com/OniHeartVT", "英語", "米国（ユーザー指定）"),
    ("Marisol Fable", "https://x.com/MarisolFable", "英語", "米国（ユーザー指定）"),
    ("DiceyDev", "https://x.com/diceydev", "英語", "米国（ユーザー指定）"),
    ("Nova Lux", "https://x.com/NovaLuxVT", "英語", "米国（ユーザー指定）"),
    ("Line Klein", "https://x.com/LineKlein100", "英語", "米国（ユーザー指定）"),
    ("Kazura", "https://x.com/KazuraArts", "英語", "米国（ユーザー指定）"),
    ("Oshen Allure", "https://x.com/OshenAllure", "英語", "米国（ユーザー指定）"),
    ("Cloverpop", "https://x.com/cloverpopvt", "英語", "米国（ユーザー指定）"),
    ("SlingerGames", "https://x.com/Slinger92", "英語", "米国（ユーザー指定）"),
    ("Fairy Mysterious Adventures", "https://x.com/FairyMystGames", "英語", "米国（ユーザー指定）"),
    ("JeansGaming", "https://x.com/JeansGamingYT", "英語", "米国（ユーザー指定）"),
    ("Kyumera", "https://x.com/KyumeraKyo", "英語", "米国（ユーザー指定）"),
    ("MizukoKanagawa", "https://x.com/MizukoKanagawa", "英語", "米国（ユーザー指定）"),
    ("Uxie", "https://x.com/Uxieuuu", "英語", "米国（ユーザー指定）"),
    ("Vivi / M_Vixn", "https://x.com/M_Vixxxnnn", "英語", "米国（ユーザー指定）"),
]


def rows_as_dicts(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row)]


def text(value):
    return str(value or "").strip()


def parse_number(value):
    if isinstance(value, (int, float)):
        return int(value)
    digits = re.sub(r"[^0-9]", "", text(value))
    return int(digits) if digits else None


def size_class(row):
    values = [parse_number(row.get("YouTube登録者")), parse_number(row.get("Twitchフォロワー"))]
    values = [value for value in values if value is not None]
    if not values:
        return "不明"
    maximum = max(values)
    if maximum <= 999:
        return "小（〜999）"
    if maximum <= 4999:
        return "小中（1,000〜4,999）"
    return "中（5,000〜49,999）"


def style_table(ws, widths=None, freeze="A2"):
    navy = PatternFill("solid", fgColor="243B53")
    white = Font(name="Yu Gothic", color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = navy
        cell.font = white
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Yu Gothic", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    if widths:
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width


def add_hyperlinks(ws, headers):
    for field in ("X URL", "YouTube URL", "Twitch URL", "証拠URL1", "証拠URL2", "証拠URL3"):
        if field not in headers:
            continue
        col = headers.index(field) + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, col)
            if text(cell.value).startswith("http"):
                cell.hyperlink = text(cell.value)
                cell.style = "Hyperlink"


def add_validations(ws, headers):
    options = {
        "送信済": '"未送信,送信済,送信不可"',
        "返信": '"未返信,返信あり,辞退"',
        "デモプレイ": '"未確認,予定,プレイ済"',
        "配信/動画化": '"未確認,予定,公開済"',
    }
    for field, formula in options.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=False)
        ws.add_data_validation(dv)
        letter = get_column_letter(headers.index(field) + 1)
        dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")
    keys = DataValidation(type="whole", operator="between", formula1="0", formula2="8", allow_blank=True)
    ws.add_data_validation(keys)
    letter = get_column_letter(headers.index("製品版キー予定人数") + 1)
    keys.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")


def write_final_sheet(book, title, rows):
    ws = book.create_sheet(title)
    ws.append(FINAL_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in FINAL_HEADERS])
    widths = [7, 8, 24, 20, 13, 22, 31, 28, 16, 34, 15, 31, 16, 42, 31, 55,
              38, 38, 31, 38, 42, 42, 42, 14, 14, 68, 12, 14, 12, 14, 14, 18, 70]
    style_table(ws, widths, "C2")
    add_hyperlinks(ws, FINAL_HEADERS)
    add_validations(ws, FINAL_HEADERS)
    ws.sheet_properties.tabColor = "70AD47" if title == "最終営業リスト" else ("F4B183" if title.startswith("A_") else "9DC3E6")
    ws["A1"].comment = Comment("A100再監査済み。Aは厳格な格上げ条件を満たす独立営業窓口。DM済み61組は除外。", "Codex")
    ws.conditional_formatting.add(
        f"AA2:AA{max(ws.max_row, 2)}",
        FormulaRule(formula=['AA2="送信済"'], fill=PatternFill("solid", fgColor="C6E0B4")),
    )
    return ws


def copy_values_sheet(source, target, source_name, target_name=None):
    old = source[source_name]
    ws = target.create_sheet(target_name or source_name)
    for row in old.iter_rows():
        for cell in row:
            new = ws[cell.coordinate]
            new.value = cell.value
            if cell.has_style:
                new.font = copy(cell.font)
                new.fill = copy(cell.fill)
                new.border = copy(cell.border)
                new.alignment = copy(cell.alignment)
                new.number_format = cell.number_format
            if cell.hyperlink:
                new._hyperlink = copy(cell.hyperlink)
    for key_, dim in old.column_dimensions.items():
        ws.column_dimensions[key_].width = dim.width
    ws.freeze_panes = old.freeze_panes
    ws.sheet_view.showGridLines = False
    return ws


def exclusion_urls(contact):
    x_url = next(iter(re.findall(r"https?://(?:www\.)?(?:x|twitter)\.com/[^\s/]+", contact)), "")
    youtube = next(iter(re.findall(r"https?://(?:www\.)?youtube\.com/[^\s]+", contact)), "")
    bilibili = next(iter(re.findall(r"https?://(?:www\.)?(?:space\.)?bilibili\.com/[^\s]+", contact)), "")
    return x_url, youtube or bilibili


def append_dm_done_exclusions(ws):
    headers = [cell.value for cell in ws[1]]
    existing_names = {normalize(ws.cell(row, 2).value) for row in range(2, ws.max_row + 1)}
    added = 0
    for name, contact, language, region in DM_DONE_RECORDS:
        if normalize(name) in existing_names:
            continue
        x_url, video_url = exclusion_urls(contact)
        values = {
            "出所": "2026-10-04 ユーザー指定DM済み", "候補名": name, "旧グループ名": "",
            "分類": "除外", "理由": "DM済み。二重営業防止のため今後の候補探索からも除外",
            "営業言語": language, "地域": region, "X URL(確認済のみ)": x_url or "未確認",
            "Twitch": "未確認", "YouTube": video_url or "未確認", "最終活動日": "未確認",
            "Twitchフォロワー": "未確認", "複数人プレイ確認数": "-", "友人/共演系確認数": "-",
            "補足": "ユーザー提示の連絡先: " + contact, "監査日": CHECKED,
        }
        ws.append([values.get(header, "") for header in headers])
        existing_names.add(normalize(name))
        added += 1
    return added


def make_progress(book, rows, dm_added):
    ws = book.create_sheet("地域別進捗")
    a_rows = [row for row in rows if row["優先度"] == "A"]
    b_rows = [row for row in rows if row["優先度"] == "B"]
    x_count = sum(text(row["X URL"]).startswith("http") for row in rows)
    unique_urls = set()
    for row in rows:
        for field in ("X URL", "YouTube URL", "Twitch URL", "証拠URL1", "証拠URL2", "証拠URL3"):
            if text(row.get(field)).startswith("http"):
                unique_urls.add(text(row[field]).lower())
    summary = [
        ["NICO DRAW 実況者・配信者営業リスト FINAL A100", "件数", "注記"],
        ["A 最優先", len(a_rows), "厳格なA再判定を通過"],
        ["B 営業対象", len(b_rows), "A条件の一部が不足"],
        ["A+B 独立営業窓口", len(rows), "最終営業リスト総数"],
        ["今回BからAへ格上げ", sum("B→A" in text(row["備考"]) for row in rows), "追加証拠・直近性・本人確認を再監査"],
        ["本人確認済みX", x_count, "Xなしは公式プロフィール掲載の別経路"],
        ["X以外の営業窓口", len(rows) - x_count, "公式外部リンク経由"],
        ["確認済みユニークURL", len(unique_urls), "動画ID付きURLへ修復済み"],
        ["ユーザー指定DM済み", len(DM_DONE_RECORDS), f"除外シート新規追記 {dm_added}件（既存除外との重複は追記せず）"],
        ["最終リスト内DM済み一致", 0, "名前・X・YouTube・Bilibiliを照合"],
        [None, None, None],
        ["営業言語", "A", "B", "合計", "送信済"],
    ]
    for item in summary:
        ws.append(item)
    for language in sorted({text(row["営業言語"]) for row in rows}):
        subset = [row for row in rows if text(row["営業言語"]) == language]
        ws.append([language, sum(row["優先度"] == "A" for row in subset), sum(row["優先度"] == "B" for row in subset), len(subset), sum(row["送信済"] == "送信済" for row in subset)])
    ws.append([None, None, None, None, None])
    ws.append(["地域", "A", "B", "合計", "送信済"])
    for region in sorted({text(row["地域"]) for row in rows}):
        subset = [row for row in rows if text(row["地域"]) == region]
        ws.append([region, sum(row["優先度"] == "A" for row in subset), sum(row["優先度"] == "B" for row in subset), len(subset), sum(row["送信済"] == "送信済" for row in subset)])
    ws.append([None, None, None, None, None])
    ws.append(["規模区分", "A", "B", "合計", "送信済"])
    for size in ("小（〜999）", "小中（1,000〜4,999）", "中（5,000〜49,999）", "不明"):
        subset = [row for row in rows if size_class(row) == size]
        ws.append([size, sum(row["優先度"] == "A" for row in subset), sum(row["優先度"] == "B" for row in subset), len(subset), sum(row["送信済"] == "送信済" for row in subset)])
    style_table(ws, [45, 14, 75, 14, 14], "A2")
    ws.sheet_properties.tabColor = "FFD966"
    return len(unique_urls)


def append_audit_rows(ws, rows, score_map):
    headers = [cell.value for cell in ws[1]]
    for row in rows:
        score = score_map[normalize(row["営業先名"])]
        proofs = score.get("proofs", [])
        values = {
            "調査区分": "2026-10-04 A100最終再監査", "候補名": row["営業先名"],
            "判定": row["優先度"], "理由": score["reason"], "営業言語": row["営業言語"],
            "X URL": row["X URL"], "YouTube URL": row["YouTube URL"], "Twitch URL": row["Twitch URL"],
            "最終活動日": row["最終活動日"], "経過日数": score["activity_age"],
            "YouTube登録者": row["YouTube登録者"], "複数人プレイ確認数": len(proofs),
            "確認ゲーム": row["確認ゲーム"], "証拠URL1": row["証拠URL1"],
            "証拠URL2": row["証拠URL2"], "証拠URL3": row["証拠URL3"], "確認日": CHECKED,
            "補足": "YouTube動画ID欠落を修復。DM済み61組との照合済み。",
        }
        ws.append([values.get(header, "") for header in headers])


def main():
    if len(DM_DONE_RECORDS) != 61 or len(DM_DONE_NAMES) != 61:
        raise RuntimeError("DM済みリストは61組である必要があります")
    source = openpyxl.load_workbook(SOURCE, data_only=False)
    score_data = json.loads(SCORED.read_text(encoding="utf-8"))["scored"]
    score_map = {normalize(item["name"]): item for item in score_data}
    current_rows = rows_as_dicts(source["最終営業リスト"])
    rebuilt = []
    for current in current_rows:
        score = score_map[normalize(current["営業先名"])]
        if score["dm_done"] or is_dm_done(current["営業先名"], current["X URL"], current["YouTube URL"], current["Twitch URL"]):
            continue
        proofs = [item for item in score.get("proofs", []) if text(item.get("url")).startswith("http")]
        proof_urls = list(dict.fromkeys(text(item["url"]) for item in proofs))[:3]
        proof_urls += [""] * (3 - len(proof_urls))
        if not proof_urls[0]:
            raise RuntimeError(f"証拠URLを復元できません: {current['営業先名']}")
        previous_priority = current["優先度"]
        current["優先度"] = "A" if score["a_qualified"] else "B"
        current["証拠URL1"], current["証拠URL2"], current["証拠URL3"] = proof_urls
        games = list(dict.fromkeys(text(item.get("game")) for item in proofs if text(item.get("game")) and "未確認" not in text(item.get("game"))))
        if games:
            current["確認ゲーム"] = " / ".join(games)
        if current["優先度"] == "A":
            current["複数人プレイ確認"] = f"A再監査済み（直接証拠{len(proof_urls) - proof_urls.count('')}件）"
            current["NICO DRAW適合理由"] = f"A判定: {score['reason']}。" + text(current["NICO DRAW適合理由"])
        else:
            current["複数人プレイ確認"] = f"確認済み（A条件不足・直接証拠{len(proof_urls) - proof_urls.count('')}件）"
            current["NICO DRAW適合理由"] = f"B判定: {score['reason']}。" + text(current["NICO DRAW適合理由"])
        current["確認日"] = CHECKED
        markers = ["A100最終再監査済", "YouTube動画ID付き証拠URLへ修復"]
        if previous_priority == "B" and current["優先度"] == "A":
            markers.insert(0, "B→A格上げ")
        current["備考"] = (text(current["備考"]) + " / " + " / ".join(markers)).strip(" /")
        rebuilt.append(current)

    rebuilt.sort(key=lambda row: (row["優先度"] != "A", normalize(row["営業先名"])))
    for index, row in enumerate(rebuilt, 1):
        row["No."] = index

    book = openpyxl.Workbook()
    book.remove(book.active)
    write_final_sheet(book, "最終営業リスト", rebuilt)
    write_final_sheet(book, "A_最優先", [row for row in rebuilt if row["優先度"] == "A"])
    write_final_sheet(book, "B_営業対象", [row for row in rebuilt if row["優先度"] == "B"])

    hold = copy_values_sheet(source, book, "保留")
    style_table(hold, [22, 26, 24, 13, 58, 15, 25, 32, 32, 35, 14, 16, 17, 18, 60, 14], "B2")
    hold.sheet_properties.tabColor = "FFD966"

    excluded = copy_values_sheet(source, book, "除外")
    dm_added = append_dm_done_exclusions(excluded)
    style_table(excluded, [22, 28, 24, 13, 64, 15, 25, 34, 32, 40, 14, 16, 17, 18, 72, 14], "B2")
    excluded.sheet_properties.tabColor = "C00000"

    group = copy_values_sheet(source, book, "グループ統合監査")
    style_table(group, [28, 28, 48, 25, 16, 42, 42, 42, 42, 16, 65], "A2")
    group.sheet_properties.tabColor = "8064A2"

    unique_urls = make_progress(book, rebuilt, dm_added)

    dm = copy_values_sheet(source, book, "DMテンプレ")
    style_table(dm, [18, 90, 90, 75, 35], "A2")
    dm.sheet_properties.tabColor = "5B9BD5"

    audit = copy_values_sheet(source, book, "調査ログ")
    append_audit_rows(audit, rebuilt, score_map)
    style_table(audit, [22, 27, 15, 62, 15, 32, 35, 32, 14, 12, 16, 18, 38, 43, 43, 43, 14, 65], "B2")
    audit.sheet_properties.tabColor = "A5A5A5"

    order = ["最終営業リスト", "A_最優先", "B_営業対象", "保留", "除外", "グループ統合監査", "地域別進捗", "DMテンプレ", "調査ログ"]
    book._sheets = [book[name] for name in order]
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.save(OUTPUT)

    counts = Counter(row["優先度"] for row in rebuilt)
    languages = Counter(row["営業言語"] for row in rebuilt if row["優先度"] == "A")
    regions = Counter(row["地域"] for row in rebuilt if row["優先度"] == "A")
    sizes = Counter(size_class(row) for row in rebuilt if row["優先度"] == "A")
    print(json.dumps({
        "output": str(OUTPUT), "A": counts["A"], "B": counts["B"], "total": len(rebuilt),
        "b_to_a": sum("B→A" in text(row["備考"]) for row in rebuilt),
        "dm_done_total": len(DM_DONE_RECORDS), "dm_done_in_final": 0, "dm_rows_added_to_excluded": dm_added,
        "hold": hold.max_row - 1, "excluded": excluded.max_row - 1, "unique_urls": unique_urls,
        "a_languages": languages, "a_regions": regions, "a_sizes": sizes,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
