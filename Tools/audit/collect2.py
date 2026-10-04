"""第2段階: Twitchは直近40アーカイブ、YouTubeはRSSで最新動画と登録者数を取得。統合メンバーも対象。"""
import re, json, time, html
import requests
import openpyxl
from concurrent.futures import ThreadPoolExecutor
import xml.etree.ElementTree as ET

SRC = r'C:\Users\mizut\Downloads\NICO_DRAW_creator_outreach_group_leaders_v200.xlsx'
OUT = r'c:\workSpace\PICO\Tools\audit\audit_raw2.json'
CID = 'kd1unb4b3q4t58fwlpcbzcbnm76a8fp'
H = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36',
     'Accept-Language': 'en-US,en;q=0.9'}


def twitch(login):
    q = ('query{user(login:"%s"){id displayName description createdAt followers{totalCount} '
         'lastBroadcast{startedAt title} channel{socialMedias{name url}} '
         'videos(first:40,type:ARCHIVE,sort:TIME){edges{node{title createdAt lengthSeconds game{name}}}}}}') % login
    err = None
    for _ in range(3):
        try:
            r = requests.post('https://gql.twitch.tv/gql', json={'query': q}, headers={'Client-ID': CID}, timeout=30)
            u = r.json().get('data', {}).get('user')
            if u is None:
                return {'login': login, 'exists': False}
            u['exists'] = True
            u['login'] = login
            u['videos'] = [e['node'] for e in (u.get('videos') or {}).get('edges', [])]
            return u
        except Exception as e:
            err = str(e)
            time.sleep(1.5)
    return {'login': login, 'error': err}


def youtube(url):
    res = {'url': url}
    try:
        r = requests.get(url, headers=H, timeout=30, cookies={'CONSENT': 'YES+1'})
        res['status'] = r.status_code
        if r.status_code != 200:
            return res
        t = r.text
        m = re.search(r'<meta property="og:title" content="([^"]+)"', t)
        res['title'] = html.unescape(m.group(1)) if m else None
        handle = re.search(r'youtube\.com/(@[^/?]+)', url)
        subs = None
        for m in re.finditer(r'⁨(@[^⁩]+)⁩ • ⁨([\d.,]+[KMB]?) subscribers⁩', t):
            if handle is None or m.group(1).lower() == requests.utils.unquote(handle.group(1)).lower():
                subs = m.group(2)
                break
        if subs is None:
            m = re.search(r'⁨([\d.,]+[KMB]?) subscribers⁩', t)
            subs = m.group(1) if m else None
        res['subs'] = subs
        cid = re.search(r'"externalId":"(UC[\w-]{22})', t) or re.search(r'channel_id=(UC[\w-]{22})', t)
        res['ytid'] = cid.group(1) if cid else None
        xs = set(requests.utils.unquote(x) for x in re.findall(r'(?:https?%3A%2F%2F)?(?:x|twitter)\.com%2F[A-Za-z0-9_]+', t))
        res['x_links'] = sorted(xs)[:10]
        if res['ytid']:
            f = requests.get('https://www.youtube.com/feeds/videos.xml?channel_id=' + res['ytid'], headers=H, timeout=30)
            vids = []
            if f.status_code == 200:
                ns = {'a': 'http://www.w3.org/2005/Atom'}
                root = ET.fromstring(f.content)
                for e in root.findall('a:entry', ns)[:15]:
                    vids.append({'title': e.find('a:title', ns).text, 'date': e.find('a:published', ns).text[:10],
                                 'url': e.find('a:link', ns).get('href')})
            res['videos'] = vids
    except Exception as e:
        res['error'] = str(e)
    return res


def main():
    wb = openpyxl.load_workbook(SRC)
    tw, yt = set(), set()
    ws = wb.worksheets[0]
    for r in ws.iter_rows(min_row=2, values_only=True):
        s = str(r[10] or '')
        for m in re.finditer(r'twitch\.tv/([A-Za-z0-9_]+)', s):
            tw.add(m.group(1).lower())
        for m in re.finditer(r'(https?://(?:www\.)?youtube\.com/(?:@[^\s;/?]+|channel/[\w-]+))', s):
            yt.add(m.group(1))
    ws = wb.worksheets[1]
    for r in ws.iter_rows(min_row=2, values_only=True):
        s = str(r[5] or '')
        for m in re.finditer(r'twitch\.tv/([A-Za-z0-9_]+)', s):
            tw.add(m.group(1).lower())
        for m in re.finditer(r'(https?://(?:www\.)?youtube\.com/(?:@[^\s;/?]+|channel/[\w-]+))', s):
            yt.add(m.group(1))
    print(len(tw), len(yt))
    with ThreadPoolExecutor(6) as ex:
        tw_res = list(ex.map(twitch, sorted(tw)))
        yt_res = list(ex.map(youtube, sorted(yt)))
    json.dump({'twitch': {r['login']: r for r in tw_res}, 'youtube': {r['url']: r for r in yt_res}},
              open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('done')


if __name__ == '__main__':
    main()
