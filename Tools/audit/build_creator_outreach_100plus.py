from __future__ import annotations

from collections import Counter
from copy import copy
from datetime import date, datetime
from pathlib import Path
import json
import re
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"C:\workSpace\PICO")
SOURCE = ROOT / "NICO_DRAW_creator_outreach_100PLUS_WORKING.xlsx"
DISCOVERY = ROOT / "Tools" / "audit" / "youtube_discovery.json"
OUTPUT = ROOT / "NICO_DRAW_creator_outreach_FINAL_100PLUS.xlsx"
CHECKED = "2026-10-03"
TODAY = date.fromisoformat(CHECKED)

FINAL_HEADERS = [
    "No.", "優先度", "営業先名", "営業単位", "営業言語", "地域",
    "X URL", "X本人確認", "Xフォロワー", "YouTube URL", "YouTube登録者",
    "Twitch URL", "Twitchフォロワー", "主な連絡方法", "DM可否",
    "代替連絡方法", "複数人プレイ確認", "共演者", "固定/準固定/単発",
    "確認ゲーム", "証拠URL1", "証拠URL2", "証拠URL3", "最終活動日",
    "確認日", "NICO DRAW適合理由", "送信済", "送信日", "返信",
    "デモプレイ", "配信/動画化", "製品版キー予定人数", "備考",
]

LANG_JA = {
    "English": "英語", "Japanese": "日本語", "Korean": "韓国語",
    "Spanish": "スペイン語", "Portuguese": "ポルトガル語",
    "Traditional Chinese": "繁体字中国語", "German": "ドイツ語",
    "Thai": "タイ語", "French": "フランス語",
}

REGION = {
    "英語": "英語圏（国未確認）", "日本語": "日本",
    "韓国語": "韓国語圏（国未確認）", "スペイン語": "スペイン語圏（国未確認）",
    "ポルトガル語": "ポルトガル語圏（国未確認）",
    "繁体字中国語": "繁体字中国語圏（国・地域未確認）",
    "ドイツ語": "ドイツ語圏（国未確認）", "タイ語": "タイ",
    "フランス語": "フランス語圏（国未確認）",
}

BAD_CHANNELS = {
    "VTClipsDayo": "切り抜き・クリップ集で、独立した実況者本人の営業窓口ではない",
    "RussieDear Clipped": "切り抜き・クリップ集で、独立した実況者本人の営業窓口ではない",
    "Holostars Central": "複数タレントの切り抜き・まとめチャンネルで、個人営業窓口ではない",
    "Aria Reacts": "ゲーム実況ではなく他者動画へのリアクション中心で、本人の複数人プレイ実績にならない",
}


def rows_as_dicts(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]


def text(value) -> str:
    return str(value or "").strip()


def clean_url(value) -> str:
    value = text(value)
    if not value.startswith("http"):
        return ""
    value = value.split("?", 1)[0].rstrip("/")
    value = value.replace("twitter.com/", "x.com/")
    value = value.replace("www.x.com/", "x.com/")
    return value


def key(value) -> str:
    return re.sub(r"[^a-z0-9\u3040-\u30ff\u3400-\u9fff\uac00-\ud7af]+", "", text(value).lower())


def bad_reason(value) -> str:
    target = key(value)
    for name, reason in BAD_CHANNELS.items():
        if key(name) == target:
            return reason
    return ""


def first_url(value) -> str:
    match = re.search(r"https?://[^\s、,]+", text(value))
    return clean_url(match.group(0)) if match else ""


def number(value):
    if isinstance(value, (int, float)):
        return int(value)
    digits = re.sub(r"[^0-9]", "", text(value))
    return int(digits) if digits else None


def age_days(value: str) -> int:
    try:
        return (TODAY - date.fromisoformat(text(value)[:10])).days
    except ValueError:
        return 99999


def has_multiplayer_signal(candidate: dict) -> bool:
    # These titles are cooperative/multiplayer by design; for the remaining games,
    # require an explicit collab/person/team signal in the concrete video title.
    intrinsic = {
        "MECCHA CHAMELEON", "PICO PARK", "PICO PARK 2", "Party Animals",
        "We Were Here", "Keep Talking and Nobody Explodes", "Pummel Party",
        "Gang Beasts", "Heave Ho", "Unrailed!",
    }
    patterns = [
        r"\bcollab(?:oration)?\b", r"\bwith\b", r"\bw/\b", r"\bfeat\.?\b", r"\bft\.?\b",
        r"\bfriends?\b", r"\btogether\b", r"\bmultiplayer\b", r"\bco-?op\b",
        r"\bteam(?:work)?\b", r"\bduo\b", r"\bsquad\b",
        r"コラボ", r"一緒", r"みんな", r"友達", r"友人", r"リア友", r"参加型",
        r"[234４３２]人", r"合流", r"共演",
        r"합방", r"같이", r"친구", r"[234]인",
        r"多人", r"一起", r"很多人", r"觀眾", r"陪我", r"朋友", r"聯動", r"联动",
        r"เพื่อน", r"ด้วยกัน", r"ทีม", r"\bw/",
        r"\bcon\b", r"\bcom\b", r"amigos?", r"juntos?", r"voc[eê]s", r"equipo", r"equipe", r"colab",
        r"\bmit\b", r"zusammen",
    ]
    for proof in candidate.get("proofs", []):
        if text(proof.get("game")) in intrinsic:
            return True
        title = text(proof.get("title")).lower()
        if any(re.search(pattern, title, re.I) for pattern in patterns):
            return True
    return False


def infer_language(candidate: dict) -> str:
    sample = " ".join(text(p.get("title")) for p in candidate.get("proofs", []))
    letters = [c for c in sample if c.isalpha()]
    total = max(len(letters), 1)
    hangul = sum("\uac00" <= c <= "\ud7af" for c in sample)
    kana = sum(("\u3040" <= c <= "\u30ff") for c in sample)
    thai = sum("\u0e00" <= c <= "\u0e7f" for c in sample)
    han = sum("\u3400" <= c <= "\u9fff" for c in sample)
    if hangul / total > 0.08:
        return "韓国語"
    if thai / total > 0.08:
        return "タイ語"
    if kana / total > 0.035:
        return "日本語"
    if han / total > 0.10 and kana / total < 0.02:
        return "繁体字中国語"
    return LANG_JA.get(candidate.get("language"), text(candidate.get("language")) or "英語")


def extract_partners(candidate: dict) -> str:
    own = key(candidate.get("name"))
    found = []
    for proof in candidate.get("proofs", []):
        title = text(proof.get("title"))
        for handle in re.findall(r"@([A-Za-z0-9_.-]{3,30})", title):
            if key(handle) != own and handle.lower() not in {x.lower() for x in found}:
                found.append(handle)
    if found:
        return " / ".join("@" + item for item in found[:8])
    return "動画タイトルで複数人プレイ確認（共演者名は未確認）"


def size_class(row: dict) -> str:
    values = [number(row.get("YouTube登録者")), number(row.get("Twitchフォロワー"))]
    values = [value for value in values if value is not None]
    if not values:
        return "不明"
    maximum = max(values)
    if maximum <= 999:
        return "小（〜999）"
    if maximum <= 4999:
        return "小中（1,000〜4,999）"
    return "中（5,000〜49,999）"


def existing_final_rows(source_book) -> list[dict]:
    rows = rows_as_dicts(source_book["最終営業リスト"])
    converted = []
    for old in rows:
        evidence_urls = [clean_url(u) for u in text(old.get("証拠URL")).splitlines() if clean_url(u)]
        proof_urls = [u for u in evidence_urls if "x.com/" not in u][:3]
        while len(proof_urls) < 3:
            proof_urls.append("")
        x_url = clean_url(old.get("X URL(本人確認済のみ)"))
        yt_url = clean_url(old.get("YouTube URL"))
        twitch_url = clean_url(old.get("Twitch URL"))
        alternative = []
        if yt_url:
            alternative.append("YouTube: " + yt_url)
        if twitch_url:
            alternative.append("Twitch: " + twitch_url)
        main_contact = text(old.get("主な連絡方法"))
        if "http" not in main_contact:
            if x_url:
                main_contact = f"{main_contact} / X: {x_url}".strip(" /")
            elif twitch_url:
                main_contact = f"{main_contact} / Twitch: {twitch_url}".strip(" /")
            elif yt_url:
                main_contact = f"{main_contact} / YouTube: {yt_url}".strip(" /")
        converted_row = {
            "No.": None,
            "優先度": old.get("優先度"),
            "営業先名": old.get("営業先名"),
            "営業単位": old.get("営業単位") or "個人（営業窓口候補）",
            "営業言語": old.get("営業言語"),
            "地域": old.get("地域"),
            "X URL": x_url or "未確認",
            "X本人確認": old.get("X本人確認") or ("確認済" if x_url else "未確認"),
            "Xフォロワー": old.get("Xフォロワー") or "未確認",
            "YouTube URL": yt_url or "未確認",
            "YouTube登録者": old.get("YouTube登録者") if old.get("YouTube登録者") not in (None, "") else "未確認",
            "Twitch URL": twitch_url or "未確認",
            "Twitchフォロワー": old.get("Twitchフォロワー") if old.get("Twitchフォロワー") not in (None, "") else "未確認",
            "主な連絡方法": main_contact,
            "DM可否": old.get("DM可否"),
            "代替連絡方法": " / ".join(alternative) or old.get("公開メール") or "要追加確認",
            "複数人プレイ確認": old.get("複数人プレイ確認"),
            "共演者": old.get("一緒に遊ぶ相手"),
            "固定/準固定/単発": old.get("固定/準固定/単発"),
            "確認ゲーム": old.get("確認ゲーム"),
            "証拠URL1": proof_urls[0], "証拠URL2": proof_urls[1], "証拠URL3": proof_urls[2],
            "最終活動日": old.get("最終活動日"), "確認日": CHECKED,
            "NICO DRAW適合理由": old.get("NICO DRAW適合理由"),
            "送信済": old.get("DM送信済") or "未送信", "送信日": "",
            "返信": old.get("返信") or "未返信", "デモプレイ": old.get("デモプレイ") or "未確認",
            "配信/動画化": old.get("配信/動画化") or "未確認",
            "製品版キー予定人数": old.get("製品版キー予定人数"),
            "備考": (text(old.get("備考")) + " / 100PLUS再監査済").strip(" /"),
            "_source": "既存42件再監査", "_size": "不明",
        }
        converted_row["_size"] = size_class(converted_row)
        converted.append(converted_row)
    return converted


def discovery_score(candidate: dict) -> tuple:
    proofs = candidate.get("proofs", [])
    return (
        len(candidate.get("games", [])), len(proofs), -age_days(candidate.get("last_active")),
        -min(number(candidate.get("subscribers")) or 999999, 999999),
    )


def candidate_to_row(candidate: dict) -> dict:
    proofs = sorted(candidate.get("proofs", []), key=lambda p: text(p.get("date")), reverse=True)[:3]
    proof_urls = [clean_url(p.get("url")) for p in proofs]
    while len(proof_urls) < 3:
        proof_urls.append("")
    games = list(dict.fromkeys(text(p.get("game")) for p in candidate.get("proofs", []) if text(p.get("game"))))
    proof_count = len(candidate.get("proofs", []))
    x_url = clean_url(candidate.get("x_url"))
    yt_url = clean_url(candidate.get("youtube_url"))
    twitch_url = clean_url(candidate.get("twitch_url"))
    external = [clean_url(u) for u in candidate.get("external_links", []) if clean_url(u)]
    non_x = [u for u in external if "x.com/" not in u and "twitter.com/" not in u]
    primary_url = x_url or (non_x[0] if non_x else "")
    if x_url:
        main_contact = "X DM: " + x_url
        dm = "X DMを試行（開放可否は送信時確認）"
        verified = "確認済（YouTube公式About欄の外部リンク）"
    else:
        domain = re.sub(r"^www\.", "", re.sub(r"^https?://", "", primary_url).split("/", 1)[0])
        main_contact = (domain + " 経由: " + primary_url) if primary_url else "要追加確認"
        dm = "外部リンク先のDM/問い合わせ手段を送信時確認"
        verified = "未確認（X掲載なし）"
    alternatives = []
    for url in non_x[1 if not x_url and non_x else 0:]:
        alternatives.append(url)
    if twitch_url:
        alternatives.append("Twitch: " + twitch_url)
    if yt_url:
        alternatives.append("YouTube: " + yt_url)
    language = infer_language(candidate)
    subs = number(candidate.get("subscribers"))
    priority = "A" if x_url and age_days(candidate.get("last_active")) <= 21 and proof_count >= 2 and (len(games) >= 2 or proof_count >= 3) and (subs or 0) <= 25000 else "B"
    frequency = "準固定候補（複数実績・固定とは断定せず）" if proof_count >= 2 else "単発確認（継続性は要追加確認）"
    fit = f"直近{age_days(candidate.get('last_active'))}日以内に{', '.join(games)}の複数人プレイ動画を{proof_count}件確認。"
    fit += f"YouTube登録者{subs:,}人の小〜中規模候補。" if subs is not None else "規模は送信時に再確認。"
    row = {
        "No.": None, "優先度": priority, "営業先名": candidate.get("name"),
        "営業単位": "個人（営業窓口候補）", "営業言語": language,
        "地域": REGION.get(language, f"{language}圏（国未確認）"),
        "X URL": x_url or "未確認", "X本人確認": verified,
        "Xフォロワー": "未確認", "YouTube URL": yt_url or "未確認",
        "YouTube登録者": subs if subs is not None else "未確認",
        "Twitch URL": twitch_url or "未確認", "Twitchフォロワー": "未確認",
        "主な連絡方法": main_contact, "DM可否": dm,
        "代替連絡方法": " / ".join(dict.fromkeys(alternatives)) or "YouTubeチャンネル経由で追加確認",
        "複数人プレイ確認": f"確認（直接動画{proof_count}件・{len(games)}ゲーム）",
        "共演者": extract_partners(candidate), "固定/準固定/単発": frequency,
        "確認ゲーム": " / ".join(games),
        "証拠URL1": proof_urls[0], "証拠URL2": proof_urls[1], "証拠URL3": proof_urls[2],
        "最終活動日": candidate.get("last_active"), "確認日": CHECKED,
        "NICO DRAW適合理由": fit, "送信済": "未送信", "送信日": "",
        "返信": "未返信", "デモプレイ": "未確認", "配信/動画化": "未確認",
        "製品版キー予定人数": None,
        "備考": "新規発掘。YouTube検索→公式About欄/RSS照合。検索語: " + text(proofs[0].get("query") if proofs else ""),
        "_source": "2026-10-03 YouTube新規調査",
    }
    row["_size"] = size_class(row)
    return row


def esbrillo_row() -> dict:
    row = {
        "No.": None, "優先度": "B", "営業先名": "EsBrilloG", "営業単位": "個人（営業窓口候補）",
        "営業言語": "スペイン語", "地域": "スペイン語圏（国未確認）", "X URL": "未確認",
        "X本人確認": "未確認（Twitch公式プロフィールにX掲載なし）", "Xフォロワー": "未確認",
        "YouTube URL": "https://www.youtube.com/@BrilloG2", "YouTube登録者": "未確認",
        "Twitch URL": "https://www.twitch.tv/esbrillog", "Twitchフォロワー": 1248,
        "主な連絡方法": "TikTok: https://www.tiktok.com/@brillog",
        "DM可否": "TikTok DMを試行（開放可否は送信時確認）",
        "代替連絡方法": "Twitch: https://www.twitch.tv/esbrillog / YouTube: https://www.youtube.com/@BrilloG2",
        "複数人プレイ確認": "確認（Twitch個別VOD 1件）",
        "共演者": "VODタイトルから特定できず", "固定/準固定/単発": "単発確認（継続性は要追加確認）",
        "確認ゲーム": "MECCHA CHAMELEON", "証拠URL1": "https://www.twitch.tv/videos/2888685239",
        "証拠URL2": "", "証拠URL3": "", "最終活動日": "2026-10-01", "確認日": CHECKED,
        "NICO DRAW適合理由": "直近2日以内にMECCHA CHAMELEONを配信。Twitch公式プロフィールからTikTokとYouTubeへの本人リンクを確認。",
        "送信済": "未送信", "送信日": "", "返信": "未返信", "デモプレイ": "未確認",
        "配信/動画化": "未確認", "製品版キー予定人数": None,
        "備考": "保留から復帰。旧記載Xは不採用のまま、Twitch公式プロフィールのTikTokを営業経路として採用。",
        "_source": "保留再監査から復帰", "_size": "小中（1,000〜4,999）",
    }
    return row


def dedupe_candidates(candidates: list[dict], existing: list[dict]):
    existing_names = {key(r["営業先名"]) for r in existing}
    existing_urls = set()
    for row in existing:
        for field in ("X URL", "YouTube URL", "Twitch URL"):
            url = clean_url(row.get(field))
            if url:
                existing_urls.add(url.lower())
    chosen = []
    rejected = []
    seen_names = set(existing_names)
    seen_urls = set(existing_urls)
    for candidate in sorted(candidates, key=discovery_score, reverse=True):
        urls = {clean_url(candidate.get(field)).lower() for field in ("x_url", "youtube_url", "twitch_url") if clean_url(candidate.get(field))}
        name_key = key(candidate.get("name"))
        if name_key in seen_names or urls & seen_urls:
            rejected.append((candidate, "既存または新規候補内で同一名/公式URLが重複"))
            continue
        chosen.append(candidate)
        seen_names.add(name_key)
        seen_urls.update(urls)
    return chosen, rejected


def style_table(ws, widths=None, freeze="A2"):
    navy = PatternFill("solid", fgColor="243B53")
    white = Font(name="Yu Gothic", color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = navy
        cell.font = white
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Yu Gothic", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    if widths:
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width


def add_hyperlinks(ws, headers):
    url_fields = {"X URL", "YouTube URL", "Twitch URL", "証拠URL1", "証拠URL2", "証拠URL3"}
    for field in url_fields:
        if field not in headers:
            continue
        col = headers.index(field) + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, col)
            if text(cell.value).startswith("http"):
                cell.hyperlink = text(cell.value)
                cell.style = "Hyperlink"


def add_validations(ws, headers):
    options = {
        "送信済": '"未送信,送信済,送信不可"',
        "返信": '"未返信,返信あり,辞退"',
        "デモプレイ": '"未確認,予定,プレイ済"',
        "配信/動画化": '"未確認,予定,公開済"',
    }
    for field, formula in options.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=False)
        dv.errorTitle = "入力値エラー"
        dv.error = "一覧から選択してください。"
        ws.add_data_validation(dv)
        letter = get_column_letter(headers.index(field) + 1)
        dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")
    keys = DataValidation(type="whole", operator="between", formula1="0", formula2="8", allow_blank=True)
    ws.add_data_validation(keys)
    letter = get_column_letter(headers.index("製品版キー予定人数") + 1)
    keys.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")
    date_col = headers.index("送信日") + 1
    for row in range(2, ws.max_row + 1):
        ws.cell(row, date_col).number_format = "yyyy-mm-dd"


def write_final_sheet(book, title, rows):
    ws = book.create_sheet(title)
    ws.append(FINAL_HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in FINAL_HEADERS])
    widths = [7, 8, 24, 20, 13, 22, 31, 28, 16, 34, 15, 31, 16, 42, 31, 55,
              30, 38, 31, 38, 42, 42, 42, 14, 14, 62, 12, 14, 12, 14, 14, 18, 60]
    style_table(ws, widths, "C2")
    add_hyperlinks(ws, FINAL_HEADERS)
    add_validations(ws, FINAL_HEADERS)
    ws.sheet_properties.tabColor = "70AD47" if title == "最終営業リスト" else ("F4B183" if title.startswith("A_") else "9DC3E6")
    ws["A1"].comment = Comment("1行=独立した営業窓口。固定グループ名だけを宛先にせず、同一URLの重複を除外済み。", "Codex")
    ws.conditional_formatting.add(
        f"AA2:AA{max(ws.max_row, 2)}",
        FormulaRule(formula=['AA2="送信済"'], fill=PatternFill("solid", fgColor="C6E0B4")),
    )
    return ws


def copy_values_sheet(source, target_book, source_name, target_name=None):
    old = source[source_name]
    ws = target_book.create_sheet(target_name or source_name)
    for row in old.iter_rows():
        for cell in row:
            new = ws[cell.coordinate]
            new.value = cell.value
            if cell.has_style:
                new.font = copy(cell.font)
                new.fill = copy(cell.fill)
                new.border = copy(cell.border)
                new.alignment = copy(cell.alignment)
                new.number_format = cell.number_format
            if cell.hyperlink:
                new._hyperlink = copy(cell.hyperlink)
    for key_, dim in old.column_dimensions.items():
        ws.column_dimensions[key_].width = dim.width
    ws.freeze_panes = old.freeze_panes
    ws.sheet_view.showGridLines = False
    return ws


def make_region_progress(book, final_rows, discovery_count, restored_count, duplicate_count, search_duplicates, hold_count, excluded_count, group_rows):
    ws = book.create_sheet("地域別進捗")
    ws.append(["NICO DRAW 実況者・配信者営業リスト FINAL 100PLUS", "件数", "注記"])
    a_count = sum(r["優先度"] == "A" for r in final_rows)
    b_count = len(final_rows) - a_count
    x_verified = sum(text(r["X URL"]).startswith("http") for r in final_rows)
    unique_urls = set()
    for row in final_rows:
        for field in ("X URL", "YouTube URL", "Twitch URL", "証拠URL1", "証拠URL2", "証拠URL3"):
            url = clean_url(row.get(field))
            if url:
                unique_urls.add(url.lower())
    ws.append(["最終営業対象（A+B）", len(final_rows), "1行=独立営業窓口"])
    ws.append(["A 最優先", a_count, "直近性・複数実績・本人確認済みXを重視"])
    ws.append(["B 営業対象", b_count, "最低条件を満たす営業可能候補"])
    ws.append(["今回の新規追加", discovery_count, "公式About欄/RSS・直接動画を確認"])
    ws.append(["保留から復帰", restored_count, "公式プロフィールから新たな営業経路を確認"])
    ws.append(["本人確認済みXあり", x_verified, "Xなし候補は公式プロフィールの別経路を使用"])
    ws.append(["X以外の営業窓口", len(final_rows) - x_verified, "YouTube/Twitch公式プロフィール掲載の別経路"])
    ws.append(["確認済みユニークURL", len(unique_urls), "公式窓口・チャンネル・直接証拠URL"])
    ws.append(["最終候補間の重複除外", duplicate_count, "同一名または同一公式URL"])
    ws.append(["検索結果の重複統合", search_duplicates, "同一チャンネルの複数検索ヒットを1営業先へ統合"])
    ws.append(["保留", hold_count, "条件不足または追加確認待ち"])
    ws.append(["除外", excluded_count, "営業対象外理由を保持"])
    ws.append(["グループ統合維持", sum(text(r.get("統合維持/解除")) == "統合維持" for r in group_rows), "監査シート参照"])
    ws.append(["グループ統合解除", sum(text(r.get("統合維持/解除")) == "統合解除" for r in group_rows), "監査シート参照"])
    ws.append([None, None, None])
    ws.append(["営業言語", "A", "B", "合計", "送信済"])
    languages = sorted({text(r["営業言語"]) for r in final_rows})
    for language in languages:
        subset = [r for r in final_rows if text(r["営業言語"]) == language]
        ws.append([language, sum(r["優先度"] == "A" for r in subset), sum(r["優先度"] == "B" for r in subset), len(subset), sum(r["送信済"] == "送信済" for r in subset)])
    ws.append([None, None, None, None, None])
    ws.append(["地域", "A", "B", "合計", "送信済"])
    regions = sorted({text(r["地域"]) for r in final_rows})
    for region in regions:
        subset = [r for r in final_rows if text(r["地域"]) == region]
        ws.append([region, sum(r["優先度"] == "A" for r in subset), sum(r["優先度"] == "B" for r in subset), len(subset), sum(r["送信済"] == "送信済" for r in subset)])
    ws.append([None, None, None, None, None])
    ws.append(["規模区分", "A", "B", "合計", "送信済"])
    sizes = ["小（〜999）", "小中（1,000〜4,999）", "中（5,000〜49,999）", "不明"]
    for size in sizes:
        subset = [r for r in final_rows if r.get("_size") == size or (size == "不明" and "不明" in text(r.get("_size")))]
        ws.append([size, sum(r["優先度"] == "A" for r in subset), sum(r["優先度"] == "B" for r in subset), len(subset), sum(r["送信済"] == "送信済" for r in subset)])
    style_table(ws, [45, 14, 68, 14, 14], "A2")
    language_header_row = 17
    ws.auto_filter.ref = f"A{language_header_row}:E{language_header_row+len(languages)}"
    ws.sheet_properties.tabColor = "FFD966"
    return len(unique_urls)


def append_hold_discovery(ws, candidate, reason):
    headers = [c.value for c in ws[1]]
    language = infer_language(candidate)
    values = {
        "出所": "2026-10-03 YouTube新規調査", "候補名": candidate.get("name"), "旧グループ名": "",
        "分類": "保留", "理由": reason, "営業言語": language, "地域": REGION.get(language, "国未確認"),
        "X URL(確認済のみ)": clean_url(candidate.get("x_url")) or "未確認",
        "Twitch": clean_url(candidate.get("twitch_url")) or "未確認",
        "YouTube": clean_url(candidate.get("youtube_url")) or "未確認",
        "最終活動日": candidate.get("last_active"), "Twitchフォロワー": "未確認",
        "複数人プレイ確認数": len(candidate.get("proofs", [])), "友人/共演系確認数": len(candidate.get("proofs", [])),
        "補足": "直接動画URL: " + " / ".join(text(p.get("url")) for p in candidate.get("proofs", [])[:3]), "監査日": CHECKED,
    }
    ws.append([values.get(h, "") for h in headers])


def append_excluded_discovery(ws, candidate, reason):
    headers = [c.value for c in ws[1]]
    language = infer_language(candidate)
    values = {
        "出所": "2026-10-03 YouTube新規調査", "候補名": candidate.get("name"), "旧グループ名": "",
        "分類": "除外", "理由": reason, "営業言語": language, "地域": REGION.get(language, "国未確認"),
        "X URL(確認済のみ)": clean_url(candidate.get("x_url")) or "未確認",
        "Twitch": clean_url(candidate.get("twitch_url")) or "未確認",
        "YouTube": clean_url(candidate.get("youtube_url")) or "未確認",
        "最終活動日": candidate.get("last_active"), "Twitchフォロワー": "未確認",
        "複数人プレイ確認数": len(candidate.get("proofs", [])), "友人/共演系確認数": len(candidate.get("proofs", [])),
        "補足": "営業先本人ではなく転載/まとめ用途のため不採用", "監査日": CHECKED,
    }
    ws.append([values.get(h, "") for h in headers])


def make_audit_log(book, source_book, candidates, selected_names, duplicate_rejected):
    ws = book.create_sheet("調査ログ")
    headers = [
        "調査区分", "候補名", "判定", "理由", "営業言語", "X URL", "YouTube URL", "Twitch URL",
        "最終活動日", "経過日数", "YouTube登録者", "複数人プレイ確認数", "確認ゲーム",
        "証拠URL1", "証拠URL2", "証拠URL3", "確認日", "補足",
    ]
    ws.append(headers)
    old = rows_as_dicts(source_book["全件監査ログ"])
    for row in old:
        evidence = [clean_url(u) for u in text(row.get("根拠")).splitlines() if clean_url(u)][:3]
        evidence += [""] * (3 - len(evidence))
        ws.append([
            "既存リスト全件再監査", row.get("候補名"), row.get("分類"), row.get("理由"), row.get("営業言語"),
            row.get("X URL"), "", "", row.get("最終活動日"), row.get("経過日数"), row.get("YouTube登録者(数値)"),
            row.get("複数人プレイ確認数"), "", evidence[0], evidence[1], evidence[2], CHECKED, row.get("補足"),
        ])
    duplicate_names = {key(c.get("name")): reason for c, reason in duplicate_rejected}
    for candidate in candidates:
        name = text(candidate.get("name"))
        if bad_reason(name):
            decision, reason = "除外", bad_reason(name)
        elif age_days(candidate.get("last_active")) > 31:
            decision, reason = "保留", "最終活動日が直近31日基準を超過"
        elif not has_multiplayer_signal(candidate):
            decision, reason = "保留", "具体的な動画タイトルだけでは複数人プレイを確定できない"
        elif key(name) in duplicate_names:
            decision, reason = "重複除外", duplicate_names[key(name)]
        elif key(name) in selected_names:
            decision, reason = "A/B採用", "活動・複数人プレイ・本人公式連絡経路を確認"
        else:
            decision, reason = "保留", "最終選定条件の再確認待ち"
        proofs = candidate.get("proofs", [])[:3]
        urls = [text(p.get("url")) for p in proofs] + [""] * (3 - len(proofs))
        ws.append([
            "YouTube新規発掘", name, decision, reason, infer_language(candidate), clean_url(candidate.get("x_url")) or "未確認",
            clean_url(candidate.get("youtube_url")) or "未確認", clean_url(candidate.get("twitch_url")) or "未確認",
            candidate.get("last_active"), age_days(candidate.get("last_active")), candidate.get("subscribers"),
            len(candidate.get("proofs", [])), " / ".join(candidate.get("games", [])), urls[0], urls[1], urls[2], CHECKED,
            "公式About欄外部リンク・公式RSS・直接動画URLを照合",
        ])
    rejected = json.loads(DISCOVERY.read_text(encoding="utf-8")).get("rejected", {})
    ws.append(["検索母集団集計", "連絡先不足", "除外集計", f"公式About欄に営業可能な外部連絡先なし: {rejected.get('no_contact', 0)}件", "", "", "", "", "", "", "", "", "", "", "", "", CHECKED, "個別採用リストには含めず"])
    ws.append(["検索母集団集計", "規模超過", "除外集計", f"小〜中規模の目安を超過: {rejected.get('too_large', 0)}件", "", "", "", "", "", "", "", "", "", "", "", "", CHECKED, "個別採用リストには含めず"])
    ws.append(["検索母集団集計", "活動基準外", "除外集計", f"直近活動基準外: {rejected.get('inactive', 0)}件", "", "", "", "", "", "", "", "", "", "", "", "", CHECKED, "個別採用リストには含めず"])
    style_table(ws, [22, 27, 15, 56, 15, 32, 35, 32, 14, 12, 16, 18, 38, 43, 43, 43, 14, 55], "B2")
    add_hyperlinks(ws, headers)
    ws.sheet_properties.tabColor = "A5A5A5"


def add_portuguese_template(ws):
    exists = any(text(ws.cell(r, 1).value) == "ポルトガル語" for r in range(2, ws.max_row + 1))
    if exists:
        return
    ws.append([
        "ポルトガル語",
        "Olá! Somos a Chaos Cat Games, uma pequena equipe indie japonesa. NICO DRAW é um jogo cooperativo de ação em que os jogadores desenham o próprio corpo e usam a quantidade de tinta como habilidade. Vimos seus vídeos de jogos cooperativos e gostaríamos de oferecer acesso à demo. Se tiver interesse, podemos enviar os detalhes e as chaves necessárias para o grupo.",
        "Olá! Gostaríamos de apresentar NICO DRAW, um jogo cooperativo indie japonês em que cada jogador desenha o próprio personagem. Caso DM não esteja disponível, qual é o melhor contato para receber informações e uma chave da demo?",
        "Olá! Enviamos uma mensagem sobre NICO DRAW. É um convite sem compromisso para testar a demo cooperativa; agradecemos se puder verificar quando tiver tempo.",
        "（公開時に差し替え）",
    ])


def add_korean_template(ws):
    exists = any(text(ws.cell(r, 1).value) == "韓国語" for r in range(2, ws.max_row + 1))
    if exists:
        return
    ws.append([
        "韓国語",
        "안녕하세요! 일본 인디 게임 팀 Chaos Cat Games입니다. NICO DRAW는 플레이어가 직접 캐릭터의 몸을 그리고, 사용한 잉크 양을 능력으로 활용하는 협동 액션 게임입니다. 협동 게임 영상을 보고 연락드렸습니다. 관심이 있으시면 데모 정보와 함께 필요한 인원수만큼 키를 보내드리고 싶습니다.",
        "안녕하세요! 직접 그린 캐릭터로 함께 플레이하는 일본 인디 협동 게임 NICO DRAW를 소개드리고 싶습니다. DM 수신이 어려운 경우, 데모 안내와 키를 받을 수 있는 연락처를 알려주시면 감사하겠습니다.",
        "안녕하세요! NICO DRAW 데모 플레이 제안 메시지를 보내드렸습니다. 부담 없는 초대이니 편하실 때 확인 부탁드립니다. 감사합니다!",
        "（公開時に差し替え）",
    ])


def main():
    source = openpyxl.load_workbook(SOURCE, data_only=False)
    discovery = json.loads(DISCOVERY.read_text(encoding="utf-8"))
    all_candidates = discovery["candidates"]
    existing = existing_final_rows(source)

    eligible = [
        c for c in all_candidates
        if not bad_reason(c.get("name"))
        and age_days(c.get("last_active")) <= 31
        and c.get("proofs")
        and has_multiplayer_signal(c)
        and (clean_url(c.get("x_url")) or any(clean_url(u) for u in c.get("external_links", [])))
    ]
    chosen, duplicate_rejected = dedupe_candidates(eligible, existing)
    new_rows = [candidate_to_row(c) for c in chosen]
    restored = esbrillo_row()
    final_rows = existing + new_rows + [restored]

    # Final guard: every row must be a distinct independent contact.
    names = [key(r["営業先名"]) for r in final_rows]
    if len(names) != len(set(names)):
        raise RuntimeError("Duplicate final contact name remains")
    final_rows.sort(key=lambda r: (r["優先度"] != "A", 0 if r["_source"].startswith("既存") else 1, -int(number(r.get("YouTube登録者")) or 0), key(r["営業先名"])))
    for idx, row in enumerate(final_rows, 1):
        row["No."] = idx

    book = openpyxl.Workbook()
    book.remove(book.active)
    write_final_sheet(book, "最終営業リスト", final_rows)
    write_final_sheet(book, "A_最優先", [r for r in final_rows if r["優先度"] == "A"])
    write_final_sheet(book, "B_営業対象", [r for r in final_rows if r["優先度"] == "B"])

    hold = copy_values_sheet(source, book, "保留")
    for row in range(hold.max_row, 1, -1):
        if text(hold.cell(row, 2).value).lower() == "esbrillog":
            hold.delete_rows(row)
    stale = [c for c in all_candidates if not bad_reason(c.get("name")) and age_days(c.get("last_active")) > 31]
    for candidate in stale:
        append_hold_discovery(hold, candidate, "最終活動日が直近31日基準を超過（再活性化時に再確認）")
    weak_multiplayer = [
        c for c in all_candidates
        if not bad_reason(c.get("name")) and age_days(c.get("last_active")) <= 31 and not has_multiplayer_signal(c)
    ]
    for candidate in weak_multiplayer:
        append_hold_discovery(hold, candidate, "具体的な動画タイトルだけでは複数人プレイを確定できない（説明欄等の追加確認待ち）")
    style_table(hold, [22, 26, 24, 13, 58, 15, 25, 32, 32, 35, 14, 16, 17, 18, 60, 14], "B2")
    hold.sheet_properties.tabColor = "FFD966"

    excluded = copy_values_sheet(source, book, "除外")
    for candidate in all_candidates:
        if bad_reason(candidate.get("name")):
            append_excluded_discovery(excluded, candidate, bad_reason(candidate.get("name")))
    for candidate, reason in duplicate_rejected:
        append_excluded_discovery(excluded, candidate, reason)
    style_table(excluded, [22, 26, 24, 13, 58, 15, 25, 32, 32, 35, 14, 16, 17, 18, 60, 14], "B2")
    excluded.sheet_properties.tabColor = "C00000"

    group = copy_values_sheet(source, book, "グループ統合監査")
    style_table(group, [28, 28, 48, 25, 16, 42, 42, 42, 42, 16, 65], "A2")
    group.sheet_properties.tabColor = "8064A2"

    group_rows = rows_as_dicts(source["グループ統合監査"])
    search_duplicates = sum(max(int(c.get("search_hit_count", 1)) - 1, 0) for c in all_candidates)
    unique_urls = make_region_progress(book, final_rows, len(new_rows), 1, len(duplicate_rejected), search_duplicates, hold.max_row - 1, excluded.max_row - 1, group_rows)

    dm = copy_values_sheet(source, book, "DMテンプレ")
    add_portuguese_template(dm)
    add_korean_template(dm)
    style_table(dm, [18, 90, 90, 75, 35], "A2")
    dm.sheet_properties.tabColor = "5B9BD5"

    selected_names = {key(c.get("name")) for c in chosen}
    make_audit_log(book, source, all_candidates, selected_names, duplicate_rejected)

    # Requested sheet order.
    order = ["最終営業リスト", "A_最優先", "B_営業対象", "保留", "除外", "グループ統合監査", "地域別進捗", "DMテンプレ", "調査ログ"]
    book._sheets = [book[name] for name in order]
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.save(OUTPUT)

    counts = Counter(r["優先度"] for r in final_rows)
    langs = Counter(r["営業言語"] for r in final_rows)
    sizes = Counter(r["_size"] for r in final_rows)
    x_verified = sum(text(r["X URL"]).startswith("http") for r in final_rows)
    print(json.dumps({
        "output": str(OUTPUT), "final": len(final_rows), "A": counts["A"], "B": counts["B"],
        "new": len(new_rows), "restored": 1, "hold": hold.max_row - 1,
        "excluded": excluded.max_row - 1, "duplicates_removed": len(duplicate_rejected), "search_duplicates_merged": search_duplicates,
        "x_verified": x_verified, "unique_confirmed_urls": unique_urls,
        "languages": langs, "sizes": sizes, "stale_added_to_hold": len(stale), "weak_multiplayer_added_to_hold": len(weak_multiplayer),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
