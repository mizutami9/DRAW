from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path
import json
import re
import sys
import zipfile

import openpyxl

from a100_config import DM_DONE_NAMES, is_dm_done, normalize
from build_creator_outreach_a100 import DM_DONE_RECORDS, FINAL_HEADERS


sys.stdout.reconfigure(encoding="utf-8")
FILE = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_FINAL_A100.xlsx")
SCORED = Path(r"C:\workSpace\PICO\Tools\audit\a100_scored.json")
SHEETS = ["最終営業リスト", "A_最優先", "B_営業対象", "保留", "除外", "グループ統合監査", "地域別進捗", "DMテンプレ", "調査ログ"]
TODAY = date(2026, 10, 4)


def rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row)]


def url_key(value):
    value = str(value or "").strip().lower().rstrip("/")
    value = value.replace("twitter.com/", "x.com/").replace("www.x.com/", "x.com/")
    return value if value.startswith("http") else ""


def main():
    problems = []
    with zipfile.ZipFile(FILE) as archive:
        bad_member = archive.testzip()
        if bad_member:
            problems.append(f"ZIP破損: {bad_member}")
    book = openpyxl.load_workbook(FILE, data_only=False)
    if book.sheetnames != SHEETS:
        problems.append(f"シート順不一致: {book.sheetnames}")
    final = book["最終営業リスト"]
    if [cell.value for cell in final[1]] != FINAL_HEADERS:
        problems.append("必須列不一致")
    data = rows(final)
    a_rows = [row for row in data if row["優先度"] == "A"]
    b_rows = [row for row in data if row["優先度"] == "B"]
    if len(a_rows) < 100:
        problems.append(f"Aが100件未満: {len(a_rows)}")
    if len(a_rows) + len(b_rows) != len(data):
        problems.append("最終リストへA/B以外が混入")
    if [row["No."] for row in data] != list(range(1, len(data) + 1)):
        problems.append("No.が連番でない")

    names = Counter(normalize(row["営業先名"]) for row in data)
    if any(count > 1 for count in names.values()):
        problems.append("営業先名重複")
    for field in ("X URL", "YouTube URL", "Twitch URL"):
        urls = Counter(url_key(row[field]) for row in data if url_key(row[field]))
        duplicates = [url for url, count in urls.items() if count > 1]
        if duplicates:
            problems.append(f"{field}重複: {duplicates[:5]}")

    score_map = {normalize(item["name"]): item for item in json.loads(SCORED.read_text(encoding="utf-8"))["scored"]}
    all_proof_urls = []
    for index, row in enumerate(data, 2):
        name = row["営業先名"]
        if is_dm_done(name, row["X URL"], row["YouTube URL"], row["Twitch URL"], row["主な連絡方法"], row["代替連絡方法"]):
            problems.append(f"行{index}: DM済みが最終リストに混入 {name}")
        score = score_map.get(normalize(name))
        if not score:
            problems.append(f"行{index}: 再監査スコアなし {name}")
        elif row["優先度"] == "A" and not score["a_qualified"]:
            problems.append(f"行{index}: A基準未達 {name}")
        proofs = [row[field] for field in ("証拠URL1", "証拠URL2", "証拠URL3") if str(row[field] or "").startswith("http")]
        if not proofs:
            problems.append(f"行{index}: 直接証拠なし {name}")
        if len(proofs) != len(set(proofs)):
            problems.append(f"行{index}: 証拠URL重複 {name}")
        for proof in proofs:
            all_proof_urls.append(proof)
            if "youtube.com" in proof and not re.search(r"youtube\.com/watch\?v=[A-Za-z0-9_-]{6,}", proof):
                problems.append(f"行{index}: YouTube動画ID欠落 {proof}")
            if "twitch.tv" in proof and not re.search(r"twitch\.tv/videos/\d+", proof):
                problems.append(f"行{index}: Twitch VOD ID欠落 {proof}")
        if str(row["確認日"]) != "2026-10-04":
            problems.append(f"行{index}: 確認日不一致")
        try:
            activity_age = (TODAY - date.fromisoformat(str(row["最終活動日"])[:10])).days
            if row["優先度"] == "A" and activity_age > 31:
                problems.append(f"行{index}: Aの活動日が31日超 {name}")
        except ValueError:
            problems.append(f"行{index}: 活動日形式不正 {name}")
    if len(all_proof_urls) != len(set(all_proof_urls)):
        problems.append("異なる営業先間で証拠URLが重複")

    a_sheet = rows(book["A_最優先"])
    b_sheet = rows(book["B_営業対象"])
    if len(a_sheet) != len(a_rows) or len(b_sheet) != len(b_rows):
        problems.append("A/B分割シート件数不一致")
    if {normalize(row["営業先名"]) for row in a_sheet + b_sheet} != set(names):
        problems.append("A/B分割シート集合不一致")

    excluded = rows(book["除外"])
    excluded_names = {normalize(row.get("候補名")) for row in excluded}
    missing_dm_done = []
    for name, contact, _, _ in DM_DONE_RECORDS:
        if normalize(name) not in excluded_names:
            # A pre-existing exclusion may use a closely related alias. Verify against
            # its stored URLs/reason before declaring it missing.
            if not any(is_dm_done(name, contact, row.get("X URL(確認済のみ)"), row.get("YouTube"), row.get("補足")) for row in excluded):
                missing_dm_done.append(name)
    if missing_dm_done:
        problems.append(f"DM済み除外記録不足: {missing_dm_done}")

    for sheet_name in ("最終営業リスト", "A_最優先", "B_営業対象"):
        sheet = book[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        for field in ("X URL", "YouTube URL", "Twitch URL", "証拠URL1", "証拠URL2", "証拠URL3"):
            col = headers.index(field) + 1
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, col)
                if str(cell.value or "").startswith("http") and not cell.hyperlink:
                    problems.append(f"{sheet_name}!{cell.coordinate}: ハイパーリンク未設定")

    result = {
        "file": str(FILE), "bytes": FILE.stat().st_size, "sheets": book.sheetnames,
        "A": len(a_rows), "B": len(b_rows), "total": len(data),
        "hold": len(rows(book["保留"])), "excluded": len(excluded),
        "dm_done_count": len(DM_DONE_RECORDS), "dm_done_in_final": sum(is_dm_done(row["営業先名"], row["X URL"], row["YouTube URL"], row["Twitch URL"]) for row in data),
        "unique_proof_urls": len(set(all_proof_urls)),
        "youtube_proofs": sum("youtube.com" in url for url in all_proof_urls),
        "twitch_proofs": sum("twitch.tv" in url for url in all_proof_urls),
        "A_languages": Counter(row["営業言語"] for row in a_rows),
        "problems": problems,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if problems:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
