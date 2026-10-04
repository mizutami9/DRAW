from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from pathlib import Path
import html as html_lib
import json
import re
import shutil
import subprocess
import sys

import openpyxl
import requests

from a100_config import detect_game, is_dm_done, multiplayer_title, normalize


sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"C:\workSpace\PICO")
SOURCE = ROOT / "NICO_DRAW_creator_outreach_FINAL_100PLUS.xlsx"
OUTPUT = ROOT / "Tools" / "audit" / "a100_b_enrichment.json"
DISCOVERY = ROOT / "Tools" / "audit" / "youtube_discovery.json"
ORIGINAL = ROOT / "NICO_DRAW_creator_outreach_100PLUS_WORKING.xlsx"
CUTOFF = date(2026, 7, 1)
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36"}


def existing_rows():
    proof_map = {}
    discovery = json.loads(DISCOVERY.read_text(encoding="utf-8"))
    for candidate in discovery.get("candidates", []):
        proof_map[normalize(candidate.get("name"))] = [
            proof.get("url") for proof in candidate.get("proofs", []) if str(proof.get("url") or "").startswith("http")
        ]
    original_book = openpyxl.load_workbook(ORIGINAL, data_only=True)
    original_sheet = original_book.worksheets[0]
    for values in original_sheet.iter_rows(min_row=2, values_only=True):
        urls = [url.strip() for url in str(values[21] or "").splitlines() if url.strip().startswith("http") and "x.com/" not in url]
        proof_map[normalize(values[2])] = urls[:3]
    book = openpyxl.load_workbook(SOURCE, data_only=True)
    sheet = book.worksheets[0]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        displayed_proofs = [str(values[idx]) for idx in (20, 21, 22) if str(values[idx] or "").startswith("http")]
        proofs = proof_map.get(normalize(values[2]), displayed_proofs)
        item = {
            "no": values[0], "priority": values[1], "name": values[2], "language": values[4],
            "x_url": values[6], "youtube_url": values[9], "twitch_url": values[11],
            "last_active": str(values[23]), "games": values[19], "proofs": proofs,
        }
        if item["priority"] == "B" and len(displayed_proofs) == 1 and str(item["youtube_url"] or "").startswith("http") and not is_dm_done(item["name"], item["x_url"], item["youtube_url"], item["twitch_url"]):
            rows.append(item)
    return rows


def channel_tab(url: str, tab: str):
    base = str(url).split("/videos", 1)[0].split("/streams", 1)[0].rstrip("/")
    command = [
        shutil.which("yt-dlp") or "yt-dlp", "--flat-playlist", "--playlist-end", "50",
        "--dump-single-json", f"{base}/{tab}",
    ]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=75)
    if result.returncode:
        return {"url": url, "tab": tab, "error": (result.stderr or "unknown yt-dlp error")[-1000:]}
    payload = json.loads(result.stdout)
    entries = []
    for entry in payload.get("entries") or []:
        video_id = entry.get("id")
        title = entry.get("title") or ""
        game = detect_game(title)
        duration = entry.get("duration") or 0
        if video_id and game and duration >= 900 and multiplayer_title(title, game):
            entries.append({
                "id": video_id, "url": f"https://www.youtube.com/watch?v={video_id}",
                "title": title, "duration": duration, "game": game,
            })
    return {
        "url": url, "tab": tab, "channel_id": payload.get("channel_id") or payload.get("id"),
        "entries": entries,
    }


def fetch_rss(channel_id: str):
    if not channel_id or not channel_id.startswith("UC"):
        return {}
    try:
        response = requests.get(f"https://www.youtube.com/feeds/videos.xml?channel_id={channel_id}", headers=UA, timeout=20)
        dates = {}
        for entry in re.findall(r"<entry>(.*?)</entry>", response.text, re.S):
            vid = re.search(r"<yt:videoId>([^<]+)</yt:videoId>", entry)
            published = re.search(r"<published>(20\d{2}-\d{2}-\d{2})", entry)
            if vid and published:
                dates[vid.group(1)] = published.group(1)
        return dates
    except Exception:
        return {}


def fetch_upload_date(video_id: str):
    try:
        response = requests.get(f"https://www.youtube.com/watch?v={video_id}", headers=UA, timeout=20)
        match = re.search(r'"uploadDate":"(20\d{2}-\d{2}-\d{2})"', response.text)
        if not match:
            match = re.search(r'<meta itemprop="uploadDate" content="(20\d{2}-\d{2}-\d{2})"', response.text)
        return match.group(1) if match else None
    except Exception:
        return None


def main():
    rows = existing_rows()
    tab_results = []
    with ThreadPoolExecutor(max_workers=6) as pool:
        futures = [pool.submit(channel_tab, row["youtube_url"], tab) for row in rows for tab in ("videos", "streams")]
        for future in as_completed(futures):
            tab_results.append(future.result())

    by_url = {}
    errors = []
    for result in tab_results:
        if result.get("error"):
            errors.append(result)
            continue
        target = by_url.setdefault(result["url"], {"channel_id": result.get("channel_id"), "entries": {}})
        target["channel_id"] = target.get("channel_id") or result.get("channel_id")
        for entry in result.get("entries") or []:
            target["entries"][entry["id"]] = entry

    rss_dates = {}
    with ThreadPoolExecutor(max_workers=12) as pool:
        future_map = {pool.submit(fetch_rss, data.get("channel_id")): url for url, data in by_url.items()}
        for future in as_completed(future_map):
            rss_dates[future_map[future]] = future.result()

    # Only request individual watch pages for matched videos absent from the official RSS.
    missing_ids = set()
    for url, data in by_url.items():
        for video_id in data["entries"]:
            if video_id not in rss_dates.get(url, {}):
                missing_ids.add(video_id)
    page_dates = {}
    with ThreadPoolExecutor(max_workers=12) as pool:
        future_map = {pool.submit(fetch_upload_date, video_id): video_id for video_id in missing_ids}
        for future in as_completed(future_map):
            page_dates[future_map[future]] = future.result()

    enriched = []
    for row in rows:
        data = by_url.get(row["youtube_url"], {"entries": {}})
        existing_ids = {url.split("v=")[-1].split("&")[0] for url in row["proofs"] if "v=" in url}
        candidates = []
        for entry in data.get("entries", {}).values():
            if entry["id"] in existing_ids:
                continue
            entry = dict(entry)
            entry["date"] = rss_dates.get(row["youtube_url"], {}).get(entry["id"]) or page_dates.get(entry["id"])
            if entry["date"] and date.fromisoformat(entry["date"]) >= CUTOFF:
                candidates.append(entry)
        candidates.sort(key=lambda item: item["date"], reverse=True)
        if candidates:
            enriched.append({**row, "additional_proofs": candidates[:2]})

    OUTPUT.write_text(json.dumps({
        "source": str(SOURCE), "checked_channels": len(rows), "tab_errors": errors,
        "enriched_count": len(enriched), "enriched": enriched,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "checked_channels": len(rows), "tab_errors": len(errors), "matched_channels": len(by_url),
        "enriched_count": len(enriched), "output": str(OUTPUT),
    }, ensure_ascii=False, indent=2))
    for row in enriched:
        proof = row["additional_proofs"][0]
        print(f"{row['no']:>3} {row['name']} | {proof['date']} {proof['game']} | {proof['title']}")


if __name__ == "__main__":
    main()
