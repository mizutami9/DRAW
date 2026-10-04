from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
PATH = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_FINAL.xlsx")
SOURCE = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_WORKING.xlsx")
EXPECTED_SHEETS = ["最終営業リスト", "集計", "グループ統合監査", "保留", "除外", "全件監査ログ", "DMテンプレ"]


def table(sheet):
    headers = [cell.value for cell in sheet[1]]
    return headers, [
        dict(zip(headers, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]


def norm_url(value):
    return str(value or "").strip().rstrip("/").lower()


errors = []
with ZipFile(PATH) as archive:
    bad_member = archive.testzip()
    if bad_member:
        errors.append(f"ZIP破損: {bad_member}")

workbook = openpyxl.load_workbook(PATH, data_only=False)
workbook_values = openpyxl.load_workbook(PATH, data_only=True)
if workbook.sheetnames != EXPECTED_SHEETS:
    errors.append(f"シート順/名称: {workbook.sheetnames}")

final_headers, final_rows = table(workbook["最終営業リスト"])
hold_headers, hold_rows = table(workbook["保留"])
excluded_headers, excluded_rows = table(workbook["除外"])
group_headers, group_rows = table(workbook["グループ統合監査"])
audit_headers, audit_rows = table(workbook["全件監査ログ"])

required_headers = [
    "No.", "優先度", "営業先名", "グループ名", "営業単位", "営業言語", "地域",
    "X URL(本人確認済のみ)", "Xフォロワー", "YouTube URL", "YouTube登録者",
    "Twitch URL", "Twitchフォロワー", "主な連絡方法", "DM可否", "公開メール",
    "複数人プレイ確認", "一緒に遊ぶ相手", "固定/準固定/単発", "確認ゲーム", "根拠",
    "証拠URL", "最終活動日", "確認日", "NICO DRAW適合理由", "DM送信済", "返信",
    "デモプレイ", "配信/動画化", "製品版キー予定人数", "備考",
]
missing_headers = [header for header in required_headers if header not in final_headers]
if missing_headers:
    errors.append(f"必須列欠落: {missing_headers}")
if len(final_rows) != 42:
    errors.append(f"最終件数: {len(final_rows)}")
if Counter(row["優先度"] for row in final_rows) != Counter({"B": 31, "A": 11}):
    errors.append("A/B件数不一致")
if len(hold_rows) != 99 or len(excluded_rows) != 47 or len(audit_rows) != 173:
    errors.append("保留/除外/監査ログ件数不一致")
if Counter(row["統合維持/解除"] for row in group_rows) != Counter({"統合解除": 28, "統合維持": 2}):
    errors.append("グループ監査件数不一致")

mandatory = [
    "優先度", "営業先名", "営業単位", "営業言語", "地域", "主な連絡方法", "DM可否",
    "公開メール", "複数人プレイ確認", "一緒に遊ぶ相手", "固定/準固定/単発", "確認ゲーム",
    "根拠", "証拠URL", "最終活動日", "確認日", "NICO DRAW適合理由",
]
for index, row in enumerate(final_rows, start=2):
    for field in mandatory:
        if row.get(field) in (None, ""):
            errors.append(f"必須値欠落 {index}:{field}")
    if row["確認日"] != "2026-10-03":
        errors.append(f"確認日不一致 {index}")
    if not str(row["複数人プレイ確認"]).startswith("確認("):
        errors.append(f"複数人プレイ未確認 {index}")
    urls = [u.strip() for u in str(row["証拠URL"]).splitlines() if u.strip()]
    if not urls or any(not u.startswith("https://") for u in urls):
        errors.append(f"証拠URL不正 {index}")
    if any(re.search(r"twitch\.tv/[^/]+/videos/?$", u, re.I) for u in urls):
        errors.append(f"汎用Twitch動画URL残存 {index}")
    try:
        active = date.fromisoformat(str(row["最終活動日"]))
        if (date(2026, 10, 3) - active).days > 31:
            errors.append(f"活動31日超 {index}:{active}")
    except ValueError:
        errors.append(f"活動日形式 {index}:{row['最終活動日']}")

seen = defaultdict(list)
for index, row in enumerate(final_rows, start=2):
    seen[("name", re.sub(r"\W", "", str(row["営業先名"]).lower()))].append(index)
    for field in ("X URL(本人確認済のみ)", "YouTube URL", "Twitch URL"):
        value = norm_url(row[field])
        if value and value not in {"未確認", "-"}:
            seen[(field, value)].append(index)
duplicates = {str(key): value for key, value in seen.items() if len(value) > 1}
if duplicates:
    errors.append(f"最終リスト重複: {duplicates}")

for label, rows, reason_field in (
    ("保留", hold_rows, "理由"),
    ("除外", excluded_rows, "理由"),
):
    for index, row in enumerate(rows, start=2):
        if not row.get(reason_field):
            errors.append(f"{label}理由欠落 {index}")

formula_cells = []
for sheet in workbook.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                formula_cells.append(f"{sheet.title}!{cell.coordinate}")
if formula_cells:
    errors.append(f"未計算数式残存: {formula_cells[:10]}")

source_wb = openpyxl.load_workbook(SOURCE, data_only=False)
source_template = [[cell.value for cell in row] for row in source_wb["DMテンプレ"].iter_rows()]
final_template = [[cell.value for cell in row] for row in workbook["DMテンプレ"].iter_rows()]
if source_template != final_template:
    errors.append("DMテンプレ内容が原本と不一致")

if len(workbook["最終営業リスト"].data_validations.dataValidation) < 5:
    errors.append("進捗入力のデータ検証不足")

summary_headers = [workbook["集計"].cell(row, 1).value for row in range(1, workbook["集計"].max_row + 1)]
for required in ("監査範囲", "未確認項目", "営業言語", "地域", "規模区分", "運用メモ"):
    if required not in summary_headers:
        errors.append(f"集計セクション欠落: {required}")

digest = hashlib.sha256(PATH.read_bytes()).hexdigest().upper()
result = {
    "ok": not errors,
    "errors": errors,
    "sha256": digest,
    "size": PATH.stat().st_size,
    "sheets": workbook.sheetnames,
    "final": len(final_rows),
    "A": sum(row["優先度"] == "A" for row in final_rows),
    "B": sum(row["優先度"] == "B" for row in final_rows),
    "hold": len(hold_rows),
    "excluded": len(excluded_rows),
    "audit_log": len(audit_rows),
    "verified_x": sum(str(row["X URL(本人確認済のみ)"]).startswith("https://x.com/") for row in final_rows),
    "x_dm": sum(str(row["DM可否"]).startswith("X DM") for row in final_rows),
    "email": sum("@" in str(row["公開メール"]) for row in final_rows),
    "language": Counter(row["営業言語"] for row in final_rows),
    "region": Counter(row["地域"] for row in final_rows),
    "size_breakdown": Counter(row["規模区分"] for row in final_rows),
    "groups": Counter(row["統合維持/解除"] for row in group_rows),
    "data_validations": len(workbook["最終営業リスト"].data_validations.dataValidation),
    "formulas": len(formula_cells),
    "duplicate_keys": len(duplicates),
}
print(json.dumps(result, ensure_ascii=False, indent=2, default=dict))
if errors:
    raise SystemExit(1)
