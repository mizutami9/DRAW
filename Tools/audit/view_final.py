import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\mizut\Downloads\NICO_DRAW_creator_outreach_FINAL.xlsx')
ws = wb['最終営業リスト']
f = open(r'c:\workSpace\PICO\Tools\audit\final_view.txt', 'w', encoding='utf-8')
for r in ws.iter_rows(min_row=2, values_only=True):
    print(r[0], r[1], r[2], '|', r[5], r[6], '|', r[7], '|fol', r[12], '|', r[16], '|', r[19], '|', str(r[20])[:150], '|', r[22], '|', r[31], file=f)
print('----HOLD', file=f)
for r in wb['保留'].iter_rows(min_row=2, values_only=True):
    print(r[1], '|', r[2], '|', str(r[4])[:110], file=f)
print('----EXCL', file=f)
for r in wb['除外'].iter_rows(min_row=2, values_only=True):
    print(r[1], '|', str(r[4])[:110], file=f)
