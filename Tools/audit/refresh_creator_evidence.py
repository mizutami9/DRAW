from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import json
import re
import sys
import time

import openpyxl
import requests

sys.stdout.reconfigure(encoding="utf-8")

SOURCE = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_WORKING.xlsx")
OUTPUT = Path(r"C:\workSpace\PICO\Tools\audit\creator_evidence_refresh.json")
CLIENT_ID = "kd1unb4b3q4t58fwlpcbzcbnm76a8fp"


def twitch_login(value: object) -> str | None:
    match = re.search(r"twitch\.tv/([A-Za-z0-9_]+)", str(value or ""), re.I)
    if not match or match.group(1).lower() in {"videos", "directory"}:
        return None
    return match.group(1).lower()


def fetch(login: str) -> dict:
    query = (
        'query{user(login:"%s"){id login displayName description followers{totalCount} '
        'lastBroadcast{startedAt title} channel{socialMedias{name url}} '
        'videos(first:40,type:ARCHIVE,sort:TIME){edges{node{id title createdAt '
        'lengthSeconds game{name}}}}}}'
    ) % login
    error = None
    for attempt in range(3):
        try:
            response = requests.post(
                "https://gql.twitch.tv/gql",
                json={"query": query},
                headers={"Client-ID": CLIENT_ID},
                timeout=30,
            )
            response.raise_for_status()
            user = response.json().get("data", {}).get("user")
            if user is None:
                return {"login": login, "exists": False}
            user["exists"] = True
            user["videos"] = [edge["node"] for edge in (user.get("videos") or {}).get("edges", [])]
            return user
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
            time.sleep(1.5 * (attempt + 1))
    return {"login": login, "error": error}


def main() -> None:
    workbook = openpyxl.load_workbook(SOURCE, data_only=False)
    logins: set[str] = set()
    for sheet_name in ("最終営業リスト", "保留", "除外"):
        sheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1]}
        twitch_col = headers.get("Twitch URL") or headers.get("Twitch")
        if not twitch_col:
            continue
        for row in range(2, sheet.max_row + 1):
            login = twitch_login(sheet.cell(row, twitch_col).value)
            if login:
                logins.add(login)

    with ThreadPoolExecutor(max_workers=6) as pool:
        results = list(pool.map(fetch, sorted(logins)))
    payload = {
        "source": str(SOURCE),
        "checked_at": "2026-10-03",
        "channels": {result.get("login", ""): result for result in results},
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    exists = sum(result.get("exists") is True for result in results)
    missing = sum(result.get("exists") is False for result in results)
    errors = [result for result in results if result.get("error")]
    videos = sum(len(result.get("videos") or []) for result in results)
    print(json.dumps({
        "requested": len(logins),
        "exists": exists,
        "missing": missing,
        "errors": len(errors),
        "videos": videos,
        "output": str(OUTPUT),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
