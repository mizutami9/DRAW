"""v200の営業先を公開情報から再取得し audit_raw.json に保存する。
Twitch: 公開GQL(プロフィール, フォロワー, SNSリンク, 直近アーカイブ)
YouTube: チャンネルページHTML(登録者数, 外部リンク, 直近動画)
"""
import re, json, sys, time, html
import requests
import openpyxl
from concurrent.futures import ThreadPoolExecutor

SRC = r'C:\Users\mizut\Downloads\NICO_DRAW_creator_outreach_group_leaders_v200.xlsx'
OUT = r'c:\workSpace\PICO\Tools\audit\audit_raw.json'
CID = 'kd1unb4b3q4t58fwlpcbzcbnm76a8fp'
UA = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36',
      'Accept-Language': 'en-US,en;q=0.9'}


def twitch(login):
    q = ('query{user(login:"%s"){id displayName description createdAt followers{totalCount} '
         'lastBroadcast{startedAt title} channel{socialMedias{name url}} '
         'videos(first:12,type:ARCHIVE,sort:TIME){edges{node{title createdAt lengthSeconds game{name}}}}}}') % login
    for _ in range(3):
        try:
            r = requests.post('https://gql.twitch.tv/gql', json={'query': q}, headers={'Client-ID': CID}, timeout=25)
            u = r.json().get('data', {}).get('user')
            if u is None:
                return {'login': login, 'exists': False}
            u['exists'] = True
            u['login'] = login
            u['videos'] = [e['node'] for e in (u.get('videos') or {}).get('edges', [])]
            return u
        except Exception as e:
            err = str(e)
            time.sleep(1.5)
    return {'login': login, 'error': err}


def youtube(url):
    res = {'url': url}
    try:
        r = requests.get(url, headers=UA, timeout=30, cookies={'CONSENT': 'YES+1'})
        t = r.text
        res['status'] = r.status_code
        m = re.search(r'"subscriberCountText":"([^"]+)"', t) or re.search(r'"subscriberCountText":\{"simpleText":"([^"]+)"', t)
        res['subs'] = m.group(1) if m else None
        m = re.search(r'<meta property="og:title" content="([^"]+)"', t)
        res['title'] = html.unescape(m.group(1)) if m else None
        m = re.search(r'<link rel="canonical" href="([^"]+)"', t)
        res['canonical'] = m.group(1) if m else None
        links = set(re.findall(r'https?%3A%2F%2F(?:x|twitter)\.com%2F[A-Za-z0-9_]+', t))
        links |= set(re.findall(r'"url":"(https?://(?:www\.)?(?:x|twitter)\.com/[A-Za-z0-9_]+)"', t))
        res['x_links'] = sorted({requests.utils.unquote(l) for l in links})
        m = re.search(r'"description":"([^"]{0,400})', t)
        res['desc'] = m.group(1) if m else None
        res['ytid'] = (re.search(r'"channelId":"(UC[\w-]{22})"', t) or [None, None])[1]
    except Exception as e:
        res['error'] = str(e)
    return res


def youtube_videos(url):
    res = {'url': url, 'videos': []}
    try:
        base = url.split('?')[0].rstrip('/')
        r = requests.get(base + '/videos', headers=UA, timeout=30, cookies={'CONSENT': 'YES+1'})
        t = r.text
        for m in re.finditer(r'"videoRenderer":\{"videoId":"([\w-]{11})".*?"title":\{"runs":\[\{"text":"(.*?)"\}\].*?"publishedTimeText":\{"simpleText":"(.*?)"\}', t):
            res['videos'].append({'id': m.group(1), 'title': m.group(2), 'published': m.group(3)})
            if len(res['videos']) >= 10:
                break
    except Exception as e:
        res['error'] = str(e)
    return res


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        rows.append(d)
    tw, yt = set(), set()
    for d in rows:
        s = ' ; '.join(str(d.get(k) or '') for k in ('配信先（YouTube / Twitch）',))
        for m in re.finditer(r'twitch\.tv/([A-Za-z0-9_]+)', s):
            tw.add(m.group(1).lower())
        for m in re.finditer(r'(https?://(?:www\.)?youtube\.com/(?:@[^\s;/?]+|channel/[\w-]+))', s):
            yt.add(m.group(1))
    print('twitch', len(tw), 'youtube', len(yt))
    with ThreadPoolExecutor(6) as ex:
        tw_res = list(ex.map(twitch, sorted(tw)))
        yt_res = list(ex.map(youtube, sorted(yt)))
        ytv_res = list(ex.map(youtube_videos, sorted(yt)))
    json.dump({'rows': rows, 'twitch': {r['login']: r for r in tw_res},
               'youtube': {r['url']: r for r in yt_res}, 'youtube_videos': {r['url']: r for r in ytv_res}},
              open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
    print('done')


if __name__ == '__main__':
    main()
