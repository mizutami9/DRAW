from __future__ import annotations

from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse
import html as html_lib
import json
import re
import shutil
import subprocess
import sys

import openpyxl
import requests

sys.stdout.reconfigure(encoding="utf-8")

SOURCE = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_100PLUS_WORKING.xlsx")
OUTPUT = Path(r"C:\workSpace\PICO\Tools\audit\youtube_discovery.json")
TODAY = date(2026, 10, 3)
CUTOFF = date(2026, 7, 1)
ACTIVE_CUTOFF = date(2026, 8, 15)
UA = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36",
    "Accept-Language": "en-US,en;q=0.9",
}

QUERIES = [
    ("English", '"MECCHA CHAMELEON" collab vtuber after:2026-07-01'),
    ("English", '"PICO PARK 2" "with friends" vtuber after:2026-07-01'),
    ("English", '"PEAK" collab vtuber after:2026-07-01'),
    ("English", '"R.E.P.O." collab vtuber after:2026-07-01'),
    ("English", '"Chained Together" collab vtuber after:2026-07-01'),
    ("English", '"Lethal Company" "with friends" vtuber after:2026-07-01'),
    ("English", '"Party Animals" collab vtuber after:2026-07-01'),
    ("English", '"We Were Here" collab vtuber after:2026-07-01'),
    ("English", '"Escape Simulator 2" collab vtuber after:2026-07-01'),
    ("English", '"Overcooked 2" collab vtuber after:2026-07-01'),
    ("Japanese", 'PICO PARK 2 コラボ VTuber after:2026-07-01'),
    ("Japanese", 'PEAK コラボ VTuber after:2026-07-01'),
    ("Japanese", 'R.E.P.O. コラボ VTuber after:2026-07-01'),
    ("Japanese", 'Chained Together コラボ VTuber after:2026-07-01'),
    ("Japanese", 'メッチャカメレオン コラボ VTuber after:2026-07-01'),
    ("Korean", '피코 파크 2 합방 버튜버 after:2026-07-01'),
    ("Korean", 'PEAK 합방 버튜버 after:2026-07-01'),
    ("Korean", 'R.E.P.O. 합방 버튜버 after:2026-07-01'),
    ("Korean", '체인드 투게더 합방 버튜버 after:2026-07-01'),
    ("Korean", '리썰 컴퍼니 합방 버튜버 after:2026-07-01'),
    ("Spanish", '"PICO PARK 2" colaboración vtuber after:2026-07-01'),
    ("Spanish", '"PEAK" amigos vtuber español after:2026-07-01'),
    ("Spanish", '"R.E.P.O." colaboración vtuber español after:2026-07-01'),
    ("Spanish", '"MECCHA CHAMELEON" collab español after:2026-07-01'),
    ("Spanish", '"Lethal Company" amigos vtuber español after:2026-07-01'),
    ("Portuguese", '"PICO PARK 2" amigos vtuber português after:2026-07-01'),
    ("Portuguese", '"PEAK" amigos vtuber brasil after:2026-07-01'),
    ("Portuguese", '"R.E.P.O." collab vtuber brasil after:2026-07-01'),
    ("Portuguese", '"MECCHA CHAMELEON" vtuber brasil after:2026-07-01'),
    ("Portuguese", '"Lethal Company" amigos vtuber brasil after:2026-07-01'),
    ("Traditional Chinese", 'PICO PARK 2 合作 台灣 VTuber after:2026-07-01'),
    ("Traditional Chinese", 'PEAK 合作 台灣 VTuber after:2026-07-01'),
    ("Traditional Chinese", 'R.E.P.O. 合作 台灣 VTuber after:2026-07-01'),
    ("Traditional Chinese", 'MECCHA CHAMELEON 台灣 VTuber after:2026-07-01'),
    ("French", '"PEAK" collab vtuber français after:2026-07-01'),
    ("French", '"R.E.P.O." collab vtuber français after:2026-07-01'),
    ("German", '"PICO PARK 2" zusammen vtuber deutsch after:2026-07-01'),
    ("German", '"PEAK" zusammen vtuber deutsch after:2026-07-01'),
    ("Thai", '"MECCHA CHAMELEON" vtuber ไทย after:2026-07-01'),
]

GAME_ALIASES = {
    "MECCHA CHAMELEON": ["meccha chameleon", "メッチャカメレオン", "めっちゃカメレオン"],
    "PICO PARK 2": ["pico park 2", "pico park", "ピコパーク", "피코 파크"],
    "PEAK": ["peak"],
    "R.E.P.O.": ["r.e.p.o.", "repo"],
    "Chained Together": ["chained together", "체인드 투게더"],
    "Lethal Company": ["lethal company", "リサルカンパニー", "리썰 컴퍼니"],
    "Party Animals": ["party animals"],
    "We Were Here": ["we were here"],
    "Escape Simulator 2": ["escape simulator"],
    "Overcooked 2": ["overcooked", "オーバークック"],
}

BLOCKED_NAMES = re.compile(r"\b(clips?|highlights?|compilation|moments|news|reaction)\b", re.I)
CONTACT_DOMAINS = ("x.com/", "twitter.com/", "bsky.app/", "instagram.com/", "linktr.ee/", "carrd.co", "beacons.ai/", "solo.to/", "lit.link/", "discord.gg/", "tiktok.com/")
HEADER_DOMAINS = CONTACT_DOMAINS + ("twitch.tv/",)


def existing_keys():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    names, urls = set(), set()
    for sheet_name in ("最終営業リスト", "保留", "除外"):
        ws = wb[sheet_name]
        headers = {cell.value: cell.column for cell in ws[1]}
        name_col = headers.get("営業先名") or headers.get("候補名")
        for row in range(2, ws.max_row + 1):
            if name_col:
                names.add(re.sub(r"\W", "", str(ws.cell(row, name_col).value or "").lower()))
            for field in ("X URL(本人確認済のみ)", "X URL(確認済のみ)", "YouTube URL", "YouTube", "Twitch URL", "Twitch"):
                col = headers.get(field)
                if col:
                    value = str(ws.cell(row, col).value or "").strip().rstrip("/").lower()
                    if value.startswith("http"):
                        urls.add(value)
    return names, urls


def search(query_info):
    language, query = query_info
    exe = shutil.which("yt-dlp")
    command = [
        exe, "--flat-playlist", "--playlist-end", "30", "--dump-single-json",
        f"ytsearch30:{query}",
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", timeout=90)
        if result.returncode:
            return {"language": language, "query": query, "error": result.stderr[-1000:]}
        payload = json.loads(result.stdout)
        return {"language": language, "query": query, "entries": payload.get("entries") or []}
    except Exception as exc:
        return {"language": language, "query": query, "error": f"{type(exc).__name__}: {exc}"}


def fetch_text(url):
    response = requests.get(url, headers=UA, cookies={"CONSENT": "YES+1"}, timeout=25)
    return response.status_code, response.text


def video_date(video_id):
    try:
        status, text = fetch_text(f"https://www.youtube.com/watch?v={video_id}")
        if status != 200:
            return None
        match = re.search(r'"uploadDate":"(20\d{2}-\d{2}-\d{2})"', text)
        if not match:
            match = re.search(r'<meta itemprop="uploadDate" content="(20\d{2}-\d{2}-\d{2})"', text)
        return match.group(1) if match else None
    except Exception:
        return None


def parse_count(value):
    if not value:
        return None
    text = str(value).replace(",", "").replace(" subscribers", "").strip()
    match = re.search(r"([\d.]+)\s*([KMB]?)", text, re.I)
    if not match:
        return None
    multiplier = {"": 1, "K": 1_000, "M": 1_000_000, "B": 1_000_000_000}[match.group(2).upper()]
    return int(float(match.group(1)) * multiplier)


def strings_in(value):
    if isinstance(value, dict):
        for item in value.values():
            yield from strings_in(item)
    elif isinstance(value, list):
        for item in value:
            yield from strings_in(item)
    elif isinstance(value, str):
        yield value


def external_link_models(value):
    if isinstance(value, dict):
        model = value.get("channelExternalLinkViewModel")
        if model:
            yield model
        for item in value.values():
            yield from external_link_models(item)
    elif isinstance(value, list):
        for item in value:
            yield from external_link_models(item)


def clean_external_url(value):
    value = html_lib.unescape(value).replace("\\u0026", "&").strip()
    if value.startswith("/") and "redirect" in value:
        value = "https://www.youtube.com" + value
    if "youtube.com/redirect" in value:
        query = parse_qs(urlparse(value).query)
        value = (query.get("q") or query.get("url") or [value])[0]
    value = unquote(value).strip().rstrip("/)")
    if not value.startswith("http") and any(value.lower().startswith(domain) for domain in HEADER_DOMAINS):
        value = "https://" + value
    return value


def channel_header_links(text):
    marker = re.search(r"(?:var ytInitialData\s*=|window\[\"ytInitialData\"\]\s*=)\s*", text)
    if not marker:
        return []
    start = text.find("{", marker.end())
    if start < 0:
        return []
    try:
        payload, _ = json.JSONDecoder().raw_decode(text[start:])
    except json.JSONDecodeError:
        return []
    links = set()
    for model in external_link_models(payload):
        for raw in strings_in(model):
            value = clean_external_url(raw)
            if any(domain in value.lower() for domain in HEADER_DOMAINS):
                links.add(value)
    return sorted(links)


def channel_info(channel_id):
    url = f"https://www.youtube.com/channel/{channel_id}"
    result = {"channel_id": channel_id, "url": url, "feed": []}
    try:
        status, text = fetch_text(url + "/about")
        result["status"] = status
        title = re.search(r'<meta property="og:title" content="([^"]+)"', text)
        result["title"] = html_lib.unescape(title.group(1)) if title else None
        subs = (
            re.search(r'"subscriberCountText":"([^"]+)"', text)
            or
            re.search(r'"subscriberCountText":\{"simpleText":"([^"]+)"', text)
            or re.search(r'"subscriberCountText":\{"accessibility":\{"accessibilityData":\{"label":"([^"]+)"', text)
            or re.search(r'([\d.,]+[KMB]?) subscribers', text, re.I)
        )
        result["subscriber_text"] = subs.group(1) if subs else None
        result["subscribers"] = parse_count(result["subscriber_text"])
        # 本人確認はチャンネルヘッダーの公式外部リンクだけを採用する。
        # 動画説明欄に含まれる共演者リンクは採用しない。
        result["external_links"] = channel_header_links(text)

        feed_status, feed_text = fetch_text(f"https://www.youtube.com/feeds/videos.xml?channel_id={channel_id}")
        if feed_status == 200:
            entries = re.findall(r"<entry>(.*?)</entry>", feed_text, re.S)
            for entry in entries[:15]:
                vid = re.search(r"<yt:videoId>([^<]+)</yt:videoId>", entry)
                title_match = re.search(r"<title>(.*?)</title>", entry, re.S)
                published = re.search(r"<published>(20\d{2}-\d{2}-\d{2})", entry)
                if vid and title_match and published:
                    result["feed"].append({
                        "id": vid.group(1),
                        "title": html_lib.unescape(title_match.group(1)),
                        "date": published.group(1),
                        "url": f"https://www.youtube.com/watch?v={vid.group(1)}",
                    })
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
    return result


def detect_game(text):
    lowered = text.lower()
    for game, aliases in GAME_ALIASES.items():
        if any(alias in lowered for alias in aliases):
            return game
    return None


def detect_language(titles, fallback):
    text = " ".join(titles)
    chars = [c for c in text if c.isalpha()]
    if not chars:
        return fallback
    if sum("\uac00" <= c <= "\ud7af" for c in chars) / len(chars) > 0.15:
        return "Korean"
    if sum("\u3040" <= c <= "\u30ff" for c in chars) / len(chars) > 0.08:
        return "Japanese"
    if sum("\u0e00" <= c <= "\u0e7f" for c in chars) / len(chars) > 0.15:
        return "Thai"
    return fallback


def main():
    existing_names, existing_urls = existing_keys()
    search_results = []
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures = [pool.submit(search, query) for query in QUERIES]
        for future in as_completed(futures):
            search_results.append(future.result())

    hits = defaultdict(list)
    video_ids = set()
    errors = []
    for result in search_results:
        if result.get("error"):
            errors.append({"query": result["query"], "error": result["error"]})
            continue
        for entry in result.get("entries") or []:
            channel_id = entry.get("channel_id")
            video_id = entry.get("id")
            title = entry.get("title") or ""
            channel = entry.get("channel") or entry.get("uploader") or ""
            if not channel_id or not video_id or not detect_game(title):
                continue
            if BLOCKED_NAMES.search(channel):
                continue
            if (entry.get("duration") or 0) < 900:
                continue
            record = {
                "video_id": video_id,
                "url": f"https://www.youtube.com/watch?v={video_id}",
                "title": title,
                "duration": entry.get("duration"),
                "channel_id": channel_id,
                "channel": channel,
                "channel_url": entry.get("uploader_url") or entry.get("channel_url") or f"https://www.youtube.com/channel/{channel_id}",
                "query_language": result["language"],
                "query": result["query"],
                "game": detect_game(title),
            }
            hits[channel_id].append(record)
            video_ids.add(video_id)

    infos = {}
    with ThreadPoolExecutor(max_workers=12) as pool:
        future_map = {pool.submit(channel_info, channel_id): channel_id for channel_id in hits}
        for future in as_completed(future_map):
            infos[future_map[future]] = future.result()

    # 検索スニペットでは投稿日が欠けるため、公式チャンネルRSSの動画IDと照合する。
    recent_hits = defaultdict(list)
    for channel_id, records in hits.items():
        feed_dates = {item["id"]: item["date"] for item in infos[channel_id].get("feed") or []}
        seen = set()
        for record in records:
            record["date"] = feed_dates.get(record["video_id"])
            if not record["date"] or date.fromisoformat(record["date"]) < CUTOFF:
                continue
            if record["video_id"] not in seen:
                recent_hits[channel_id].append(record)
                seen.add(record["video_id"])

    candidates = []
    rejected = Counter()
    for channel_id, records in recent_hits.items():
        info = infos[channel_id]
        channel_name = info.get("title") or records[0]["channel"]
        normalized_name = re.sub(r"\W", "", channel_name.lower())
        channel_urls = {
            f"https://www.youtube.com/channel/{channel_id}".lower(),
            str(records[0].get("channel_url") or "").strip().rstrip("/").lower(),
        }
        if normalized_name in existing_names or any(url in existing_urls for url in channel_urls if url):
            rejected["existing"] += 1
            continue
        latest = max([item["date"] for item in info.get("feed") or []] + [item["date"] for item in records])
        if date.fromisoformat(latest) < ACTIVE_CUTOFF:
            rejected["inactive"] += 1
            continue
        subscribers = info.get("subscribers")
        if subscribers is not None and subscribers > 50_000:
            rejected["too_large"] += 1
            continue
        external = info.get("external_links") or []
        x_links = [link for link in external if "x.com/" in link.lower() or "twitter.com/" in link.lower()]
        contact_links = [link for link in external if any(domain in link.lower() for domain in CONTACT_DOMAINS)]
        if not contact_links:
            rejected["no_contact"] += 1
            continue
        feed_proofs = []
        for item in info.get("feed") or []:
            game = detect_game(item["title"])
            if game and date.fromisoformat(item["date"]) >= CUTOFF:
                feed_proofs.append({**item, "game": game})
        proofs = []
        for item in sorted(records + feed_proofs, key=lambda x: x["date"], reverse=True):
            if item["url"] not in {proof["url"] for proof in proofs}:
                proofs.append(item)
            if len(proofs) >= 3:
                break
        languages = [record["query_language"] for record in records]
        language = detect_language([record["title"] for record in records] + [item["title"] for item in info.get("feed") or []], Counter(languages).most_common(1)[0][0])
        candidates.append({
            "channel_id": channel_id,
            "name": channel_name,
            "youtube_url": records[0].get("channel_url") or info["url"],
            "subscribers": subscribers,
            "subscriber_text": info.get("subscriber_text"),
            "x_url": x_links[0] if x_links else None,
            "external_links": contact_links,
            "twitch_url": next((link for link in external if "twitch.tv/" in link.lower()), None),
            "language": language,
            "last_active": latest,
            "proofs": proofs,
            "games": sorted({proof["game"] for proof in proofs}),
            "search_hit_count": len(records),
        })

    candidates.sort(key=lambda item: (not bool(item["x_url"]), item["subscribers"] is None, item["subscribers"] or 999_999, item["name"].lower()))
    OUTPUT.write_text(json.dumps({
        "generated": TODAY.isoformat(),
        "query_count": len(QUERIES),
        "search_errors": errors,
        "raw_channels": len(hits),
        "recent_channels": len(recent_hits),
        "rejected": rejected,
        "candidate_count": len(candidates),
        "candidates": candidates,
    }, ensure_ascii=False, indent=2, default=dict), encoding="utf-8")
    print(json.dumps({
        "queries": len(QUERIES),
        "errors": len(errors),
        "raw_channels": len(hits),
        "recent_channels": len(recent_hits),
        "candidate_count": len(candidates),
        "with_x": sum(bool(candidate["x_url"]) for candidate in candidates),
        "language": Counter(candidate["language"] for candidate in candidates),
        "rejected": rejected,
        "output": str(OUTPUT),
    }, ensure_ascii=False, indent=2, default=dict))


if __name__ == "__main__":
    main()
