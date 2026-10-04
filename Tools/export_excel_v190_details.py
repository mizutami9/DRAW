import openpyxl
import json

excel_path = r'C:\Users\mizut\Downloads\NICO_DRAW_creator_outreach_group_leaders_v190.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

data = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    sheet_data = []
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        row = {}
        for c, h in enumerate(headers, 1):
            if h is not None:
                row[str(h)] = ws.cell(r, c).value
        # only keep if row has any non-empty value
        if any(v is not None for v in row.values()):
            sheet_data.append(row)
    data[sheet_name] = sheet_data

with open('Tools/excel_v190_dump.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print("Exported v190 dump successfully!")
