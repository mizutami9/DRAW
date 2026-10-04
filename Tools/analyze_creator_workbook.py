from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlparse
import json
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
PATH = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_WORKING.xlsx")
wb = openpyxl.load_workbook(PATH, data_only=False)


def rows_as_dicts(sheet_name: str):
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


final = rows_as_dicts("最終営業リスト")
hold = rows_as_dicts("保留")
excluded = rows_as_dicts("除外")
groups = rows_as_dicts("グループ統合監査")
logs = rows_as_dicts("全件監査ログ")

def norm(v):
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower())

url_fields = ["X URL(本人確認済のみ)", "YouTube URL", "Twitch URL"]
dupes = []
seen = defaultdict(list)
for sheet, rows in (("final", final), ("hold", hold), ("excluded", excluded)):
    for i, r in enumerate(rows, 2):
        for field in url_fields:
            value = r.get(field)
            if value and value not in ("未確認", "-"):
                key = (field, str(value).strip().rstrip("/").lower())
                seen[key].append((sheet, i, r.get("営業先名") or r.get("候補名") or r.get("名前")))
for key, locs in seen.items():
    if len(locs) > 1:
        dupes.append({"key": key, "locations": locs})

missing = defaultdict(list)
required = [
    "優先度", "営業先名", "営業単位", "営業言語", "地域",
    "主な連絡方法", "DM可否", "複数人プレイ確認", "固定/準固定/単発",
    "確認ゲーム", "根拠", "証拠URL", "最終活動日", "確認日", "NICO DRAW適合理由",
]
for i, r in enumerate(final, 2):
    for field in required:
        if r.get(field) in (None, ""):
            missing[field].append((i, r.get("営業先名")))

report = {
    "counts": {
        "final": len(final),
        "A": sum(r.get("優先度") == "A" for r in final),
        "B": sum(r.get("優先度") == "B" for r in final),
        "hold": len(hold),
        "excluded": len(excluded),
        "groups": len(groups),
        "logs": len(logs),
    },
    "final_language": Counter(r.get("営業言語") for r in final),
    "final_region": Counter(r.get("地域") for r in final),
    "final_size": Counter(r.get("規模区分") for r in final),
    "x_verified": Counter(r.get("X本人確認") for r in final),
    "dm": Counter(r.get("DM可否") for r in final),
    "email_available": sum(bool(r.get("公開メール") and r.get("公開メール") != "未確認") for r in final),
    "group_decisions": Counter(r.get("統合維持/解除") for r in groups),
    "missing_required": missing,
    "duplicate_urls": dupes,
    "final_rows": [
        {
            "row": i,
            "priority": r.get("優先度"),
            "name": r.get("営業先名"),
            "group": r.get("グループ名"),
            "unit": r.get("営業単位"),
            "language": r.get("営業言語"),
            "x": r.get("X URL(本人確認済のみ)"),
            "youtube": r.get("YouTube URL"),
            "twitch": r.get("Twitch URL"),
            "contact": r.get("主な連絡方法"),
            "dm": r.get("DM可否"),
            "email": r.get("公開メール"),
            "multi": r.get("複数人プレイ確認"),
            "relation": r.get("固定/準固定/単発"),
            "games": r.get("確認ゲーム"),
            "evidence_urls": r.get("証拠URL"),
            "last_active": r.get("最終活動日"),
            "x_verified": r.get("X本人確認"),
            "size": r.get("規模区分"),
        }
        for i, r in enumerate(final, 2)
    ],
}

print(json.dumps(report, ensure_ascii=False, indent=2, default=list))
