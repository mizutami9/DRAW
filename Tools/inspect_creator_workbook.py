from pathlib import Path
import json
import openpyxl
import sys

sys.stdout.reconfigure(encoding="utf-8")


path = Path(r"C:\workSpace\PICO\NICO_DRAW_creator_outreach_WORKING.xlsx")
wb = openpyxl.load_workbook(path, data_only=False)

result = {
    "sheets": [],
    "defined_names": [str(x) for x in wb.defined_names.values()],
}

for ws in wb.worksheets:
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    preview = []
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 8), values_only=True):
        preview.append(list(row))
    result["sheets"].append({
        "title": ws.title,
        "state": ws.sheet_state,
        "rows": ws.max_row,
        "cols": ws.max_column,
        "headers": headers,
        "preview": preview,
        "freeze": str(ws.freeze_panes) if ws.freeze_panes else None,
        "auto_filter": ws.auto_filter.ref,
        "merged": [str(r) for r in ws.merged_cells.ranges],
    })

print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
