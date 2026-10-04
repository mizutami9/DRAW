import openpyxl
import json
import os

excel_path = r'C:\Users\mizut\Downloads\NICO_DRAW_creator_outreach_group_leaders_v190.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

report = {}
sheet_names = wb.sheetnames

for name in sheet_names:
    ws = wb[name]
    rows = []
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in row_vals):
            rows.append(row_vals)
    report[name] = {
        'max_row': ws.max_row,
        'max_col': ws.max_column,
        'non_empty_rows': len(rows),
        'header': rows[0] if rows else [],
        'sample_row_1': rows[1] if len(rows) > 1 else [],
        'all_rows_count': len(rows)
    }

output_path = r'Tools/excel_structure_report.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump({'sheet_names': sheet_names, 'sheets': report}, f, ensure_ascii=False, indent=2)

print(f"Successfully analyzed {len(sheet_names)} sheets and saved report to {output_path}")
