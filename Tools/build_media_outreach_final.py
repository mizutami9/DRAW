from __future__ import annotations

import argparse
import copy
import hashlib
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from build_media_outreach_v9 import (
    BLUE,
    GRAY,
    GREEN,
    INK,
    NAVY,
    ORANGE,
    PALE_BLUE,
    RED,
    VERIFIED_ON,
    WHITE,
    YELLOW,
    add_title,
    base_border,
    collect_media_rows,
    create_dashboard,
    fill,
    normalize_workbook,
    set_hyperlink,
    setup_page,
    style_header,
    style_table,
)


FINAL_MEDIA = [
    [
        63,
        "Noisy Pixel",
        "北米/英語圏",
        "英語",
        "https://noisypixel.net/",
        "pr@noisypixel.net（公式Bluesky / Staff）",
        "メール",
        "2026年9月に小規模PC作品『Aokaki』の無料Steamデモを単独掲載。公式Staffは隠れたインディー作品の発掘を掲げる。",
        "https://noisypixel.net/akupara-kallisti-games-aokaki-free-steam-demo/",
        "日本発・独特なビジュアル・無料Steamデモとの相性が高い。描いた形とINK配分が能力になる点をGIFで見せる。",
        "未送信", "未確認", None, "未送信", "未送信",
        "窓口根拠: https://noisypixel.net/staff/ / 公式Blueskyプロフィール。",
        "高", "第1陣", "英語本文＋15秒GIF＋Steam＋Press Kit",
    ],
    [
        64,
        "MonsterVine",
        "北米/英語圏",
        "英語",
        "https://monstervine.com/",
        "will@monstervine.com",
        "メール",
        "2026年8月に小規模作品『Hope in the City』の無料Steamデモを単独掲載。月次インディーまとめで開発者へ直接情報提供を呼びかける。",
        "https://monstervine.com/2026/08/hope-in-the-city-demo-steam-release/",
        "デモと小規模インディーの直接提案を歓迎。協力プレイの短いクリップと、Steamですぐ試せる事実を先頭へ。",
        "未送信", "未確認", None, "未送信", "未送信",
        "窓口根拠: https://monstervine.com/2026/01/january-2026-indie-game-wrap-up/",
        "高", "第1陣", "担当者名Willを明記し個別送信",
    ],
    [
        65,
        "Hey Poor Player",
        "北米/英語圏",
        "英語",
        "https://www.heypoorplayer.com/",
        "deadpixels@heypoorplayer.com（公式連絡先）",
        "メール",
        "2026年2月のSteam Next Festで実際に遊んだ10本を選定し、開発者からの直接連絡をきっかけに含めた作品も明記。",
        "https://www.heypoorplayer.com/2026/02/24/10-unmissable-demos-during-february-2026-steam-next-fest/",
        "直接提案が選定につながった実績が明確。記事の形式に合わせ、プレイ後に語れる3点を箇条書きで示す。",
        "未送信", "未確認", None, "未送信", "未送信",
        "窓口根拠: https://www.heypoorplayer.com/privacy-policy/",
        "高", "第1陣", "Steamデモの見どころ3点を本文に記載",
    ],
    [
        66,
        "GameLuster",
        "英語圏",
        "英語",
        "https://gameluster.com/",
        "https://gameluster.com/the-team/",
        "公式Teamページ",
        "公式Teamページでプレスリリースとレビューコードの送付を案内。2026年も小規模な無料Steam作品を継続掲載。",
        "https://gameluster.com/free-steam-games-new-releases/",
        "速報より製品版レビュー/ニュース向け。第2陣で、キャラクター作成と協力パズルの独自性を短く提案する。",
        "未送信", "未確認", None, "未送信", "未送信",
        "公開アドレスは送信直前に公式Teamページで確認し、保護表示なら同ページの導線を使う。",
        "中", "第2陣", "公式Teamページの最新担当を確認",
    ],
    [
        67,
        "The Indie Informer",
        "北米/英語圏",
        "英語",
        "https://theindieinformer.com/",
        "https://theindieinformer.com/contact/",
        "Contact form",
        "インディー専門のOpenCritic/Metacritic掲載媒体。2026年8月にもデモイベントの実プレイ動画を掲載し、編集長は手描きインディーを好むと公式Aboutに明記。",
        "https://theindieinformer.com/blog/",
        "NICO DRAWの手描き表現とコアフックが媒体嗜好へ直結。定型リリースより短い個別ピッチを優先。",
        "未送信", "未確認", None, "未送信", "未送信",
        "媒体特性根拠: https://theindieinformer.com/about/",
        "高", "第1陣", "手描き表現を示すGIFと個別の一文を付ける",
    ],
    [
        68,
        "Six One Indie",
        "北米/英語圏",
        "英語",
        "https://www.sixoneindie.com/",
        "hello@sixoneindie.com（公式About）",
        "メール",
        "インディー専門のOpenCritic掲載媒体。2026年は多数のDemo Previewに加え、150本超のDemo Daysを実施。",
        "https://www.sixoneindie.com/",
        "デモ、パズル、マルチプレイの各枠に合う。編集提案と次回Showcase/Demo Daysの応募確認を分けて管理する。",
        "未送信", "未確認", None, "未送信", "未送信",
        "窓口根拠: https://www.sixoneindie.com/about/ / イベント根拠: https://www.sixoneindie.com/demo-days",
        "高", "第1陣", "編集提案として送付。イベント応募は別枠で管理",
    ],
    [
        69,
        "Checkpoint Gaming",
        "豪州/英語圏",
        "英語",
        "https://checkpointgaming.net/",
        "https://checkpointgaming.net/contact-us/",
        "Contact form",
        "2025年に協力パズル/プラットフォーム作品を詳細レビューし、2026年もPCと小規模インディーのレビューを継続。",
        "https://checkpointgaming.net/reviews/2025/03/split-fiction-review/",
        "豪州時間帯の英語媒体。協力して起きる笑える失敗と、複数解法のパズル設計を軸に第2陣で提案。",
        "未送信", "未確認", None, "未送信", "未送信",
        "公式Contact formを使用。一般コメント欄へ投稿しない。",
        "中", "第2陣", "豪州東部の平日午前に送信",
    ],
    [
        70,
        "게임포커스 (GameFocus)",
        "韓国",
        "韓国語",
        "https://www.gamefocus.co.kr/",
        "bodo@gamefocus.co.kr（公式フッターの記事提供）",
        "メール",
        "2025年6月に最大4人協力のSteam Next Festデモを単独掲載。2026年もデモを含むインディー作品を継続的に扱う。",
        "https://www.gamefocus.co.kr/detail.php?number=166254",
        "4人オンライン協力、Steamデモ、韓国語対応が掲載実績と直結。韓国語全文で送る。",
        "未送信", "未確認", None, "未送信", "未送信",
        "窓口根拠: https://www.gamefocus.co.kr/ （公式フッター）",
        "高", "第1陣", "韓国語本文＋4人協力GIF＋Steam",
    ],
    [
        71,
        "게임인사이트 (GameInsight)",
        "韓国",
        "韓国語",
        "https://www.gameinsight.co.kr/",
        "press@gameinsight.co.kr（公式記事記者連絡先）",
        "メール",
        "2025年2月のSteam Next Fest注目作を企画記事で扱い、インディーを含むデモの体験内容を紹介。",
        "https://www.gameinsight.co.kr/news/articleView.html?idxno=34696",
        "Next Fest特集候補。代表窓口の公開状態を送信日に再確認し、韓国語で短く送る。",
        "未送信", "未確認", None, "未送信", "未送信",
        "代表アドレスの根拠は同ドメイン公式記事。媒体トップの接続状態も送信前に確認。",
        "中", "第2陣", "窓口と担当分野を送信日に再確認",
    ],
    [
        72,
        "VidaoPantalla",
        "スペイン/スペイン語",
        "スペイン語",
        "https://www.vidaopantalla.es/",
        "prensa@vidaopantalla.es（公式Contacto）",
        "メール",
        "2026年も多数の小規模インディーShowcaseとSteamデモを継続整理。公式Contactoはニュースリリース専用編集部メールを公開。",
        "https://www.vidaopantalla.es/2026/06/resumen-completo-summer-game-fest-2026/",
        "日本発の風変わりなインディーをスペイン語圏へ届ける追加先。見出し化できる1文とGIFを添える。",
        "未送信", "未確認", None, "未送信", "未送信",
        "窓口根拠: https://www.vidaopantalla.es/contacto/",
        "高", "第1陣", "スペイン語全文＋短尺GIF＋Steam",
    ],
    [
        73,
        "Arkade",
        "ブラジル",
        "ブラジルポルトガル語",
        "https://arkade.com.br/",
        "https://arkade.com.br/contato/",
        "Contact form",
        "2026年2月に最大4人で遊べる小規模Steam Next Festデモを単独掲載。公式Contact formが利用可能。",
        "https://arkade.com.br/demo-gratuita-de-witchspire-ja-esta-disponivel-no-steam-teste-a-sobrevivencia-magica-com-seus-amigos/",
        "4人協力デモの掲載実績がそのまま合う。PT-BRで、友人をつかむ・投げる場面の短尺動画を付ける。",
        "未送信", "未確認", None, "未送信", "未送信",
        "フォーム送信後、送信内容と日時を営業履歴へ保存。",
        "高", "第1陣", "PT-BR本文＋協力プレイ動画",
    ],
    [
        74,
        "GameLobby",
        "ブラジル",
        "ブラジルポルトガル語",
        "https://gamelobby.com.br/",
        "contato@gamelobby.com.br（公式Press Kit指南）",
        "メール",
        "2026年8月にも小規模ブラジル産Steam作品の公開Playtestを単独掲載。公式Press Kit指南はデモURLとレビュー連絡先を推奨。",
        "https://gamelobby.com.br/cururu-whispers-of-the-forest-abre-playtest-gratuito-na-steam/",
        "小規模開発とテスト段階の作品を扱う。日本発である点、PT-BR対応、4人オンラインを冒頭で示す。",
        "未送信", "未確認", None, "未送信", "未送信",
        "窓口根拠: https://gamelobby.com.br/wp-content/uploads/2025/07/Manual-Presskit-Gamelobby-PT-BR.pdf",
        "高", "第1陣", "PT-BR本文＋Steam＋Press Kit",
    ],
]


MAILTO = {
    "Noisy Pixel": "pr@noisypixel.net",
    "MonsterVine": "will@monstervine.com",
    "Hey Poor Player": "deadpixels@heypoorplayer.com",
    "Six One Indie": "hello@sixoneindie.com",
    "게임포커스 (GameFocus)": "bodo@gamefocus.co.kr",
    "게임인사이트 (GameInsight)": "press@gameinsight.co.kr",
    "VidaoPantalla": "prensa@vidaopantalla.es",
    "GameLobby": "contato@gamelobby.com.br",
}


AUDIT_ROWS = [
    ["採用基準", "公式窓口", "必須", "公式サイト、公式Contact、公式アカウント、または媒体自身の記事で確認できる送付経路", "送信当日に再確認", ""],
    ["採用基準", "直近実績", "必須", "原則2025〜2026年の小規模インディー、Steamデモ、協力/パズルの掲載実績", "根拠記事URLを本表へ記録", ""],
    ["採用基準", "NICO DRAW適合", "必須", "手描き、キャラ作成、協力、パズル、Steamデモのいずれかが媒体テーマと接続", "定型一斉送信を避ける", ""],
    ["今回採用", "Noisy Pixel〜GameLobby", "本表No.63〜74", "公式窓口と直近掲載の両方を確認できた12媒体を追加", "第1陣9 / 第2陣3", "https://noisypixel.net/staff/"],
    ["窓口修正", "4Gamers", "採用継続", "公式YouTube概要にニュース協力先 news@4gamers.com.tw が掲載されていることを確認", "送信日に概要欄を再確認", "https://www.youtube.com/channel/UCb6QDBuxtoIP8kT4cO5BX0A/about"],
    ["別枠", "Steam Next Fest Press Preview", "PR経路へ追加", "登録情報の連絡先を選択メディアへ共有するSteam公式導線。直接営業と併用", "参加登録時に連絡先提供を確認", "https://partner.steamgames.com/doc/marketing/upcoming_events/nextfest?l=english"],
    ["別枠", "Wholesome Games", "PR経路へ追加", "イベント応募は期間制で、Web編集媒体の送信件数とは分離", "次回公募時に世界観適合を再評価", "https://wholesomegames.com/"],
    ["別枠", "Six One Indie Showcase / Demo Days", "PR経路へ追加", "編集部ピッチとイベント応募を別タスクとして管理", "次回募集開始時に応募条件を確認", "https://www.sixoneindie.com/demo-days"],
    ["保留", "Turn Based Lovers", "本表No.29", "デモ発掘力は高いがターン制専門でジャンル不一致", "ターン制要素を追加した場合のみ再検討", "https://turnbasedlovers.com/"],
    ["保留", "RPGFan", "本表No.41", "デモ特集実績はあるがRPG/アドベンチャー中心でジャンル不一致", "RPG要素を前面に出せる場合のみ再検討", "https://www.rpgfan.com/"],
    ["見送り", "遊戲基地 Gamebase", "現時点で不採用", "直近デモ記事は強いが、営業用途の正式な編集情報提供窓口を確認できず、客服窓口は使用しない", "公式の新聞投稿窓口が公開されたら再調査", "https://news.gamebase.com.tw/news/detail/99440169"],
    ["見送り", "MeriStation", "現時点で不採用", "2026年のNext Fest掲載は確認したが、営業用の公式編集窓口を確認できず", "公式Contactでプレス受付が明示されたら再調査", "https://as.com/meristation/noticias/lo-mejor-del-steam-next-fest-10-demos-que-no-te-puedes-perder-f202603-n/"],
    ["見送り", "Adrenaline", "現時点で不採用", "2026年のSteam/インディー掲載は活発だが、今回の監査では編集提案の公式送付先を確認できず", "ニュース提供専用窓口が公開されたら再調査", "https://www.adrenaline.com.br/games/convergence-games-showcase-2026/"],
    ["見送り", "大手総合媒体の追加", "現時点で不採用", "知名度だけで増やさず、直近の小規模デモ実績と正式窓口が揃う媒体を優先", "デモ反応・掲載実績が出た後に製品版向けで再検討", ""],
    ["完了定義", "FINALスナップショット", "2026-10-03", "対象言語圏を横断し、採用・保留・見送り・別枠の理由を記録。世界中の媒体が将来増えないという意味ではない", "新規追加は上記3基準を満たす場合のみ", ""],
]


PR_ROUTES = [
    [
        "Steam Next Fest Press Preview",
        "Steam公式メディア試遊導線",
        "全世界",
        "https://partner.steamgames.com/doc/marketing/upcoming_events/nextfest?l=english",
        "登録時に開発者が連絡先提供へ同意すると、選定された報道関係者へ連絡情報が共有される公式仕組み",
        "Next Fest登録、公開ストア、公開Demo、連絡先を期限前に整える",
        "必須",
        "Press Previewでは公開中のストア/Demo情報が見られる。Playtestだけに依存しない",
    ],
    [
        "Wholesome Games",
        "インディーShowcase / Steamイベント",
        "海外/英語",
        "https://wholesomegames.com/",
        "2026年も公募制ShowcaseとSteamイベントを運営し、公式サイトで今後の公募確認を案内",
        "次回公募時に、明るい手描き表現と協力の楽しさが企画趣旨に合う場合だけ応募",
        "公募待ち",
        "常設の編集記事営業先ではない。締切と条件を公式で確認",
    ],
    [
        "Six One Indie Showcase / Demo Days",
        "インディーShowcase / Demoイベント",
        "海外/英語",
        "https://www.sixoneindie.com/demo-days",
        "2026年Demo Daysで150本超のデモを扱い、マルチプレイ/協力・パズル枠を設置",
        "編集部への通常ピッチとは別に、次回募集が開いた時だけ応募",
        "公募待ち",
        "応募条件、生成AIポリシー、締切を応募回ごとに確認",
    ],
]


def add_final_media(wb) -> None:
    ws = wb["メディア営業リスト"]
    last = max(r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value)
    template_row = last
    for values in FINAL_MEDIA:
        last += 1
        for col in range(1, 20):
            src = ws.cell(template_row, col)
            dst = ws.cell(last, col)
            if src.has_style:
                dst._style = copy.copy(src._style)
            dst.alignment = copy.copy(src.alignment)
            dst.number_format = src.number_format
        for col, value in enumerate(values, start=1):
            ws.cell(last, col, value)
        ws.row_dimensions[last].height = 88
        set_hyperlink(ws.cell(last, 5), values[4])
        set_hyperlink(ws.cell(last, 9), values[8])
        if values[1] in MAILTO:
            ws.cell(last, 6).hyperlink = f"mailto:{MAILTO[values[1]]}"
            ws.cell(last, 6).font = Font(name="Yu Gothic", size=10, color="0563C1", underline="single")

    final_last = last
    ws.auto_filter.ref = f"A1:S{final_last}"
    ws.conditional_formatting.add(
        f"R2:R{final_last}", FormulaRule(formula=['$R2="第1陣"'], fill=fill(GREEN))
    )
    ws.conditional_formatting.add(
        f"R2:R{final_last}", FormulaRule(formula=['$R2="第2陣"'], fill=fill(YELLOW))
    )
    ws.conditional_formatting.add(
        f"R2:R{final_last}", FormulaRule(formula=['$R2="保留"'], fill=fill(GRAY))
    )

    # Existing Taiwan route is now verified through the outlet's official YouTube About page.
    ws.cell(18, 6, "news@4gamers.com.tw（公式YouTube概要）")
    ws.cell(18, 7, "メール")
    ws.cell(18, 16, "窓口根拠: https://www.youtube.com/channel/UCb6QDBuxtoIP8kT4cO5BX0A/about")
    ws.cell(18, 19, "ニュース協力先を公式YouTube概要で送信日に再確認")
    ws.cell(18, 6).hyperlink = "mailto:news@4gamers.com.tw"
    ws.cell(18, 6).font = Font(name="Yu Gothic", size=10, color="0563C1", underline="single")


def patch_dashboard(wb) -> None:
    rows = collect_media_rows(wb)
    create_dashboard(wb, rows)
    ws = wb["営業ダッシュボード"]
    ws["A1"] = "NICO DRAW メディア営業ダッシュボード — FINAL 2026-10-03"
    ws["A2"] = "最終監査日: 2026-10-03　採用基準・見送り理由・再調査条件は「最終監査ログ」に記録。送信直前の窓口再確認は必要。"

    ws["A17"] = "最優先で送る16媒体"
    priorities = [
        (1, "The Indie Informer", "手描きインディー嗜好＋2026年デモ実演", "英語"),
        (2, "Six One Indie", "インディー専門＋Demo Days/協力枠", "英語"),
        (3, "MonsterVine", "開発者へ直接情報提供を呼びかけ", "英語"),
        (4, "Noisy Pixel", "日本発・小規模Steamデモの直近掲載", "英語"),
        (5, "GamingTrend", "Steamデモ公開→Next Festを単独掲載", "英語"),
        (6, "GameSpew", "小規模開発のNext Festデモ＋公式受付", "英語"),
        (7, "Co-Optimus", "協力ゲーム専門。4人オンラインに直結", "英語"),
        (8, "Alpha Beta Gamer", "無料公開デモを公式Submissionで受付", "英語"),
        (9, "게임포커스", "4人協力Steamデモ＋公式記事提供メール", "韓国語"),
        (10, "게임조선", "小規模Steamデモ＋編集部メール", "韓国語"),
        (11, "인디게임닷컴", "インディー専門＋開発者広報受付", "韓国語"),
        (12, "Gamereactor 中文版", "協力パズルデモ＋公式編集窓口", "繁体字"),
        (13, "Vandal", "オンライン協力パズルのNext Fest記事", "スペイン語"),
        (14, "VidaoPantalla", "2026年インディー/デモ継続掲載", "スペイン語"),
        (15, "Arkade", "4人協力Steamデモの直近掲載", "PT-BR"),
        (16, "GameLobby", "小規模Steamテスト段階の掲載＋公式窓口", "PT-BR"),
    ]
    # Clear the previous 12-entry block and create two columns of eight.
    for row in range(18, 27):
        for col in range(1, 9):
            ws.cell(row, col).value = None
            ws.cell(row, col).fill = fill(WHITE)
            ws.cell(row, col).border = base_border()
    headers = ["順", "媒体", "選定理由", "言語", "順", "媒体", "選定理由", "言語"]
    for col, value in enumerate(headers, 1):
        ws.cell(18, col, value)
    style_header(ws, 18, 1, 8)
    for idx, item in enumerate(priorities):
        block = 0 if idx < 8 else 4
        row = 19 + (idx % 8)
        for offset, value in enumerate(item):
            cell = ws.cell(row, 1 + block + offset, value)
            cell.border = base_border()
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 42
    ws.freeze_panes = "A3"


def patch_schedule(wb) -> None:
    ws = wb["送信スケジュール"]
    # First-wave total is now 44; send in three manageable batches.
    ws.cell(3, 7, 14)
    ws.cell(4, 7, 14)
    ws.cell(5, 7, 16)
    ws.cell(5, 8, "第1陣44件送信完了")
    ws.cell(6, 7, "最大16")
    ws.cell(7, 7, 13)
    ws.cell(8, 7, 13)


def patch_research_queue(wb) -> None:
    ws = wb["追加調査キュー"]
    updates = {
        2: ("監査完了", "日本語13件。既存の高適合媒体を優先し、知名度だけの追加は行わない"),
        3: ("監査完了", "繁体字4件。4Gamersのニュース協力メールを確認。Gamebaseは正式編集窓口未確認で見送り"),
        4: ("監査完了", "韓国語8件。GameFocus / GameInsightを追加"),
        5: ("監査完了", "英語37件。Noisy Pixel / MonsterVine / Hey Poor Player / GameLuster / The Indie Informer / Six One Indie / Checkpoint Gamingを追加"),
        6: ("監査完了", "スペイン語5件。VidaoPantallaを追加。MeriStationは正式編集窓口未確認で見送り"),
        7: ("監査完了", "PT-BR 6件。Arkade / GameLobbyを追加。Adrenalineは正式編集窓口未確認で見送り"),
        8: ("不採用（現時点）", "客服窓口は営業に使わない。公式のニュース投稿/編集窓口が公開された場合だけ再調査"),
        9: ("採用済み", "VidaoPantallaをNo.72へ追加"),
        10: ("採用済み", "Arkade / GameLobbyをNo.73〜74へ追加"),
    }
    for row, (state, memo) in updates.items():
        ws.cell(row, 6, state)
        ws.cell(row, 7, memo)
    ws["C1"] = "当初目標"
    ws["F1"] = "FINAL判定"


def patch_findings(wb) -> None:
    ws = wb["調査知見"]
    additions = [
        ["FINALの意味", "2026-10-03時点で、対象地域・採用基準に沿う営業先を監査し、採用・保留・見送り・別枠を記録した確定スナップショット。世界中の媒体が将来増えないという意味ではない。"],
        ["数より適合", "Web編集媒体72件で運用上十分。これ以上は窓口未確認・ジャンル不一致・大手総合媒体への薄い一斉送信が増えるため、追加件数を完了条件にしない。"],
        ["イベント導線", "Steam Press Preview、Wholesome Games、Six One Indieの公募は通常の編集営業と分離し、応募期間と条件を公式で確認する。"],
    ]
    start = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value) + 1
    template = start - 1
    for vals in additions:
        for col, value in enumerate(vals, 1):
            cell = ws.cell(start, col, value)
            cell._style = copy.copy(ws.cell(template, col)._style)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[start].height = 48
        start += 1
    ws.auto_filter.ref = f"A1:B{start - 1}"


def add_pr_routes(wb) -> None:
    ws = wb["配信_PR経路"]
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value)
    template = 2
    for vals in PR_ROUTES:
        last += 1
        for col, value in enumerate(vals, 1):
            cell = ws.cell(last, col, value)
            cell._style = copy.copy(ws.cell(template, col)._style)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        set_hyperlink(ws.cell(last, 4), vals[3])
        ws.row_dimensions[last].height = 74
    ws.auto_filter.ref = f"A1:H{last}"
    ws.freeze_panes = "A2"


def create_audit_sheet(wb) -> None:
    if "最終監査ログ" in wb.sheetnames:
        del wb["最終監査ログ"]
    ws = wb.create_sheet("最終監査ログ", 3)
    add_title(ws, "NICO DRAW メディア営業 — 最終監査ログ", 6)
    ws.append(["区分", "対象", "判定", "根拠 / 判断", "再調査条件 / 次アクション", "証拠URL"])
    for row in AUDIT_ROWS:
        ws.append(row)
    style_header(ws, 2, 1, 6)
    style_table(ws, 3, ws.max_row, 1, 6)
    for row in range(3, ws.max_row + 1):
        category = ws.cell(row, 1).value
        color = {
            "採用基準": BLUE,
            "今回採用": GREEN,
            "窓口修正": GREEN,
            "別枠": PALE_BLUE,
            "保留": GRAY,
            "見送り": ORANGE,
            "完了定義": NAVY,
        }.get(category, WHITE)
        ws.cell(row, 1).fill = fill(color)
        ws.cell(row, 1).font = Font(
            name="Yu Gothic", size=10, bold=True,
            color=WHITE if category == "完了定義" else NAVY,
        )
        set_hyperlink(ws.cell(row, 6), ws.cell(row, 6).value)
        ws.row_dimensions[row].height = 56
    for col, width in enumerate([16, 30, 22, 78, 48, 64], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = f"A2:F{ws.max_row}"
    setup_page(ws)
    ws.freeze_panes = "A3"


def rebuild_progress(wb) -> None:
    rows = collect_media_rows(wb)
    orders = Counter(r[17] for r in rows)
    langs = Counter(r[3] for r in rows)
    ws = wb["進捗"]
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
            cell.fill = fill(WHITE)
            cell.border = base_border()
    data = [
        ["NICO DRAW メディア営業 FINAL — 進捗と監査結果", None],
        ["指標", "確定値"],
        ["調査済み営業候補（別枠含む）", len(rows)],
        ["Web編集媒体", len(rows) - orders.get("別枠", 0)],
        ["第1陣", orders.get("第1陣", 0)],
        ["第2陣", orders.get("第2陣", 0)],
        ["保留", orders.get("保留", 0)],
        ["別枠（PR配信/動画）", orders.get("別枠", 0)],
        [None, None],
        ["営業言語", "件数"],
        *[[lang, langs.get(lang, 0)] for lang in ["日本語", "英語", "韓国語", "繁体字中国語", "簡体字中国語", "スペイン語", "ブラジルポルトガル語"]],
        [None, None],
        ["FINAL追加（No.63〜74）", "採用根拠"],
        ["Noisy Pixel", "日本発/小規模Steamデモ＋公式PR窓口"],
        ["MonsterVine", "開発者へ直接情報提供を呼びかけ＋無料Steamデモ"],
        ["Hey Poor Player", "Next Fest実プレイ選定＋開発者連絡が掲載に寄与"],
        ["GameLuster", "公式Teamがリリース/レビューコード受付"],
        ["The Indie Informer", "手描きインディー嗜好＋2026年デモ実演"],
        ["Six One Indie", "インディー専門＋Demo Days/協力枠"],
        ["Checkpoint Gaming", "協力パズル実績＋公式Contact"],
        ["GameFocus", "4人協力Steamデモ＋公式記事提供メール"],
        ["GameInsight", "Next Fest企画＋公式ドメインのpress窓口"],
        ["VidaoPantalla", "2026インディー/デモ掲載＋プレス専用メール"],
        ["Arkade", "4人協力Steamデモ＋公式Contact form"],
        ["GameLobby", "小規模Steamテスト掲載＋公式Press Kit指南"],
        [None, None],
        ["FINALで閉じたもの", "内容"],
        ["採用/見送り基準", "公式窓口・直近実績・NICO DRAW適合の3条件"],
        ["候補の判断履歴", "最終監査ログに採用、保留、見送り、再調査条件を記録"],
        ["PR/Showcase導線", "Steam Press Preview / Wholesome Games / Six One Indieを別枠追加"],
        ["4Gamers窓口", "公式YouTube概要のnews@4gamers.com.twへ修正"],
        [None, None],
        ["送信前のP0", "判断"],
        ["SteamのDemo表示", "一般ユーザー環境でDemoボタンを確認。Playtestだけなら営業本文でDemo公開と断定しない"],
        ["Demoオンライン協力", "別アカウント・別回線の2台以上で接続確認"],
        ["Press Kit URL/返信先", "公開URLと監視可能な返信先を確定"],
        ["Next Fest回/日付", "Steam公式で確定した情報だけ本文へ記載"],
    ]
    for row_idx, values in enumerate(data, 1):
        ws.cell(row_idx, 1, values[0])
        ws.cell(row_idx, 2, values[1])
    ws.merge_cells("A1:B1")
    ws["A1"] = "NICO DRAW メディア営業 FINAL — 進捗と監査結果"
    ws["A1"].fill = fill(NAVY)
    ws["A1"].font = Font(name="Yu Gothic", size=15, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    section_rows = [2, 10, 19, 33, 39]
    for row in section_rows:
        style_header(ws, row, 1, 2)
    for row in range(2, len(data) + 1):
        if row not in section_rows:
            for col in range(1, 3):
                ws.cell(row, col).font = Font(name="Yu Gothic", size=10, color=INK)
                ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
                ws.cell(row, col).border = base_border()
            ws.row_dimensions[row].height = 36
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 100
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"


def validate_final(wb) -> None:
    rows = collect_media_rows(wb)
    if len(rows) != 74:
        raise ValueError(f"Expected 74 media rows, found {len(rows)}")
    if [r[0] for r in rows] != list(range(1, 75)):
        raise ValueError("Media No. values are not sequential 1..74")
    names = [r[1] for r in rows]
    if len(names) != len(set(names)):
        raise ValueError("Duplicate media names detected")
    if any(not r[8] for r in rows):
        raise ValueError("Every media row must have an evidence URL")
    if any(r[17] not in {"第1陣", "第2陣", "保留", "別枠"} for r in rows):
        raise ValueError("Unexpected send-order value")
    if "最終監査ログ" not in wb.sheetnames:
        raise ValueError("Final audit sheet is missing")
    if wb["送信テンプレート"].max_row != 22:
        raise ValueError("Expected 21 localized message templates")
    expected = Counter({"第1陣": 44, "第2陣": 26, "保留": 2, "別枠": 2})
    actual = Counter(r[17] for r in rows)
    if actual != expected:
        raise ValueError(f"Unexpected wave counts: {actual}")


def build(source: Path, output: Path) -> None:
    wb = openpyxl.load_workbook(source)
    add_final_media(wb)
    patch_dashboard(wb)
    patch_schedule(wb)
    patch_research_queue(wb)
    patch_findings(wb)
    add_pr_routes(wb)
    create_audit_sheet(wb)
    rebuild_progress(wb)

    preferred = [
        "営業ダッシュボード",
        "方針_送信タイミング",
        "メディア営業リスト",
        "最終監査ログ",
        "送信テンプレート",
        "送信スケジュール",
        "送信前チェック",
        "プレスキット準備",
        "プレスリリース骨子",
        "配信_PR経路",
        "追加調査キュー",
        "調査知見",
        "進捗",
    ]
    wb._sheets = [wb[name] for name in preferred]
    normalize_workbook(wb)
    validate_final(wb)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
    print(f"OUTPUT={args.output}")
    print(f"SIZE={args.output.stat().st_size}")
    print(f"SHA256={sha256(args.output)}")


if __name__ == "__main__":
    main()
